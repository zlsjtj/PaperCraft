"""Build a short Chinese editorial report from authored evidence and actual images.

Usage: python scripts/build_editorial_report.py report.json report.docx
This formats human/model-authored comparisons; it does not judge their truth.
Relative image paths are resolved against report.json. Outputs must be new.
"""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn


def build(source, output):
    if output.exists():
        raise ValueError('Output exists; preserve the prior report')
    spec = json.loads(source.read_text(encoding='utf-8-sig'))
    for key in ('title', 'summary', 'baseline', 'changes', 'checks', 'remaining'):
        if key not in spec:
            raise ValueError('Missing authored report field: ' + key)
    images = []
    for row in spec.get('figures', []):
        pair = [source.parent / row[key] for key in ('before', 'after')]
        if not all(p.is_file() for p in pair):
            raise ValueError('Missing actual comparison image')
        images.append((row, pair))
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(18)
    s.left_margin = s.right_margin = Mm(23)
    for name, size in [('Normal', 10), ('Title', 19), ('Heading 1', 13)]:
        st = d.styles[name]
        st.font.name, st.font.size = 'Arial', Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.paragraph_format.space_after = Pt(6)
    d.styles['Normal'].paragraph_format.line_spacing = 1.1
    d.add_paragraph(spec['title'], 'Title')
    d.add_paragraph(spec['summary'])
    d.add_paragraph('基线与交付', 'Heading 1')
    d.add_paragraph(spec['baseline'])
    for value in spec.get('final_files', []):
        d.add_paragraph(value)
    for row in spec['changes']:
        d.add_paragraph(row['focus'], 'Heading 1')
        for label, key in [('原文', 'before'), ('改文', 'after'), ('理由', 'reason'),
                           ('依据', 'evidence'), ('代价与边界', 'limits')]:
            if row.get(key):
                d.add_paragraph(label + '：' + row[key])
    for row, pair in images:
        d.add_page_break()
        d.add_paragraph(row['title'], 'Heading 1')
        for label, path in zip(('修改前', '修改后'), pair):
            d.add_paragraph(label)
            d.add_picture(str(path), width=Mm(160))
        d.add_paragraph(row['caption'])
    d.add_paragraph('验证与未解决事项', 'Heading 1')
    for row in spec['checks']:
        d.add_paragraph(row['scope'] + '：' + row['status'] + '。' + row['detail'])
    for value in spec['remaining']:
        d.add_paragraph(value)
    d.add_paragraph('本工具只排版所提供记录；技术核对、模型效果审阅、真人及作者认可分别记载。黄色标记不等于 Word 原生修订。')
    output.parent.mkdir(parents=True, exist_ok=True)
    d.save(output)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('spec', type=Path)
    p.add_argument('output', type=Path)
    a = p.parse_args()
    try:
        build(a.spec, a.output)
    except (ValueError, KeyError, OSError) as exc:
        p.exit(1, str(exc) + '\n')


if __name__ == '__main__':
    main()
