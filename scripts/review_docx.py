"""Evidence-linked plain-paragraph edits; preserves the rest of the DOCX ZIP.
This helper applies authored prose. It is not a scientific fact checker or LLM.
"""
from __future__ import annotations
import argparse,copy,hashlib,json,re,tempfile,zipfile
from pathlib import Path
from lxml import etree as E
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement,parse_xml
from docx.oxml.ns import qn

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
NS={'w':W,'m':M};Q=lambda n:'{'+W+'}'+n
PROTECTED=['m:oMath','w:tbl','w:drawing','w:pict','w:fldSimple','w:fldChar','w:instrText',
           'w:footnoteReference','w:endnoteReference','w:bookmarkStart','w:bookmarkEnd','w:hyperlink']
def sha(b):return hashlib.sha256(b).hexdigest()
def text(p):return ''.join(p.xpath('.//w:t/text() | .//m:t/text()',namespaces=NS))
def canonical(p):return E.tostring(p,method='c14n')
def read_package(path):
    with zipfile.ZipFile(path) as z:
        if z.testzip():raise ValueError('Invalid DOCX ZIP')
        infos=z.infolist();parts={i.filename:z.read(i.filename) for i in infos}
    return infos,parts,E.fromstring(parts['word/document.xml'])
def paragraphs(root):return root.find('w:body',NS).findall('w:p',NS)
def inventory(path):
    _,parts,root=read_package(path)
    return {'file':str(Path(path).resolve()),'sha256':sha(Path(path).read_bytes()),
      'paragraphs':[{'id':f'P{i:04d}','text':text(p),'text_sha256':sha(text(p).encode()),
                     'editable':plain_reason(p) is None,'reason':plain_reason(p)}
                    for i,p in enumerate(paragraphs(root),1)],
      'objects':{n:len(root.xpath('//'+n,namespaces=NS)) for n in PROTECTED},
      'media_parts':sum(n.startswith('word/media/') for n in parts)}
def plain_reason(p):
    # A refusal is preferable to flattening a field, equation, hyperlink or run styling.
    if any(c.tag not in (Q('pPr'),Q('r')) for c in p):return 'complex paragraph child'
    if p.find('w:pPr/w:sectPr',NS) is not None:return 'section boundary'
    props=[]
    for r in p.findall('w:r',NS):
        if any(c.tag not in (Q('rPr'),Q('t')) for c in r):return 'complex run content'
        rp=r.find('w:rPr',NS)
        if rp is not None and rp.find('w:highlight',NS) is not None:return 'existing highlight'
        props.append(E.tostring(rp) if rp is not None else b'')
    if len(set(props))>1:return 'mixed run formatting'
    return None
def write_text(p,value,highlight):
    if not isinstance(value,str) or not value.strip():raise ValueError('Replacement text must be nonempty')
    rp=p.find('w:r/w:rPr',NS);rp=copy.deepcopy(rp) if rp is not None else E.Element(Q('rPr'))
    for c in list(p):
        if c.tag!=Q('pPr'):p.remove(c)
    if highlight:
        rp=parse_xml(E.tostring(rp))
        rp.get_or_add_highlight().set(Q('val'),'yellow')
    r=E.SubElement(p,Q('r'));r.append(rp);t=E.SubElement(r,Q('t'))
    t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=value
def apply_plan(source,plan,highlight=True):
    root=copy.deepcopy(source);base=paragraphs(root)
    lookup={f'P{i:04d}':p for i,p in enumerate(base,1)}
    changes=[];ids=set();used=set()
    if root.xpath('//w:ins | //w:del | //w:moveFrom | //w:moveTo | //w:pPrChange | //w:rPrChange',namespaces=NS):
        raise ValueError('Existing tracked revisions: use the documents workflow without accepting them silently')
    for op in plan['operations']:
        cid=op['id'];kind=op['kind'];pid=op['target']
        if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_-]*',cid) or cid in ids:raise ValueError('Invalid/duplicate change ID')
        ids.add(cid)
        if pid not in lookup or pid in used:raise ValueError('Unknown or repeated target; combine edits per paragraph')
        p=lookup[pid]
        if p.getparent() is None:raise ValueError('Target already deleted')
        before=text(p)
        if op['expect']!=before:raise ValueError('Expected source text differs for '+pid)
        reason=plain_reason(p)
        if reason:raise ValueError(f'{pid}: {reason}; use a format-aware manual edit')
        if not op.get('purpose') or not op.get('evidence') or not isinstance(op.get('confirm'),bool):
            raise ValueError('Each change requires purpose, evidence and Boolean confirm')
        pp=p.find('w:pPr',NS);before_format=E.tostring(pp,encoding='unicode') if pp is not None else ''
        target=p
        if kind=='replace':write_text(p,op['text'],highlight)
        elif kind=='insert_after':
            target=E.Element(Q('p'))
            if pp is not None:target.append(copy.deepcopy(pp))
            # New body prose inherits the source paragraph style, not its text/runs.
            write_text(target,op['text'],highlight);p.addnext(target)
        elif kind=='delete':p.getparent().remove(p);target=None
        elif kind=='move_after':
            anchor=lookup.get(op.get('after'))
            if anchor is None or anchor is p or anchor.getparent() is None:raise ValueError('Invalid move anchor')
            anchor.addnext(p)
        elif kind=='format':
            options=op['format']
            if not options or set(options)-{'space_after_pt','keep_next'}:raise ValueError('Unsupported format property')
            if pp is None:pp=parse_xml(f'<w:pPr xmlns:w="{W}"/>');p.insert(0,pp)
            else:
                ordered=parse_xml(E.tostring(pp));p.replace(pp,ordered);pp=ordered
            for key,val in options.items():
                tag='spacing' if key=='space_after_pt' else 'keepNext'
                el=pp.get_or_add_spacing() if tag=='spacing' else pp.get_or_add_keepNext()
                if key=='space_after_pt':
                    if not isinstance(val,(int,float)) or not 0<=val<=72:raise ValueError('Invalid spacing')
                    el.set(Q('after'),str(round(val*20)))
                else:
                    if not isinstance(val,bool):raise ValueError('keep_next must be Boolean')
                    el.set(Q('val'),'1' if val else '0')
        else:raise ValueError('Unsupported operation '+kind)
        used.add(pid)
        after_pp=target.find('w:pPr',NS) if target is not None else None
        changes.append({'id':cid,'old_location':pid,'new_location':None,'type':kind,
          'before':before,'after':text(target) if target is not None else '',
          'purpose':op['purpose'],'evidence':op['evidence'],'author_confirmation_needed':op['confirm'],
          'format_before':before_format if kind=='format' else None,
          'format_after':E.tostring(after_pp,encoding='unicode') if kind=='format' and after_pp is not None else None,
          'format_requested':op.get('format') if kind=='format' else None,
          '_node':target})
    final=paragraphs(root)
    for c in changes:
        node=c.pop('_node');c['new_location']=f'P{final.index(node)+1:04d}' if node is not None else 'deleted'
    for tag in PROTECTED:
        if list(map(canonical,source.xpath('//'+tag,namespaces=NS)))!=list(map(canonical,root.xpath('//'+tag,namespaces=NS))):
            raise ValueError('Protected objects changed: '+tag)
    return root,changes
def write_package(path,infos,parts,root):
    with zipfile.ZipFile(path,'w') as z:
        for info in infos:
            payload=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True) if info.filename=='word/document.xml' else parts[info.filename]
            z.writestr(info,payload)
def verify_pair(source,output):
    _,a,ar=read_package(source);_,b,br=read_package(output)
    if set(a)!=set(b):raise ValueError('ZIP parts differ')
    changed=[n for n in a if a[n]!=b[n]]
    if set(changed)-{'word/document.xml'}:raise ValueError('Unexpected package mutation')
    for tag in PROTECTED:
        if list(map(canonical,ar.xpath('//'+tag,namespaces=NS)))!=list(map(canonical,br.xpath('//'+tag,namespaces=NS))):raise ValueError('Protected object mismatch: '+tag)
    return {'status':'PASS','changed_parts':changed,'objects':inventory(output)['objects'],
            'source_sha256':sha(Path(source).read_bytes()),'output_sha256':sha(Path(output).read_bytes()),
            'scope':'ZIP/XML object preservation only; content, citations and visual review are separate'}
def make_report(path,source,plan,changes,checks):
    report=plan['report']
    required={'summary','thesis','titles','dimensions','ledger','workload','figures','limitations','author_questions','checks'}
    if required-set(report):raise ValueError('Missing report fields: '+str(required-set(report)))
    d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.8)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=None
        st._element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if st._element.get_or_add_rPr().find(qn('w:rFonts')) is None else None
        st._element.rPr.find(qn('w:rFonts')).set(qn('w:eastAsia'),'Microsoft YaHei')
        st.paragraph_format.space_after=Pt(6)
    d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.15
    d.styles['Title'].font.size=Pt(19);d.styles['Heading 1'].font.size=Pt(13);d.styles['Heading 2'].font.size=Pt(11)
    # Eliminate inherited Word title rules/borders.
    for st in ['Title','Heading 1','Heading 2']:
        for el in d.styles[st]._element.xpath('.//w:pBdr'):el.getparent().remove(el)
    d.add_paragraph('论文包装诊断报告','Title');d.add_paragraph(report['summary'])
    def section(title):d.add_paragraph(title,'Heading 1')
    def line(value):d.add_paragraph(str(value))
    section('基线与范围');line('输入：'+Path(source).name);line('SHA-256：'+plan['input_sha256'])
    line('范围：'+plan.get('scope','full')+'；执行方式：'+plan.get('execution','sequential'))
    section('主贡献与标题');line(report['thesis'])
    for x in report['titles']:line(x['title']+'；'+x['reason'])
    section('三维修改对照')
    for x in report['dimensions']:line(x['name']+'：原稿 '+x['before']+'；本轮 '+x['after'])
    section('贡献与证据')
    for x in report['ledger']:
        line(x['id']+' '+x['claim']);line('差异：'+x['difference']+'；来源：'+x['source'])
        line('性质：'+x['nature']+'；状态：'+x['status']+'；范围：'+x['scope'])
    section('已完成工作与新增分析')
    for x in report['workload']:line(x['name']+'：'+x['detail'])
    section('图文实际状态')
    for x in report['figures']:line(x['label']+'：'+x['status']+'；'+x['detail'])
    section('主要改动与高亮位置')
    line('段落号按 document.xml 的直接正文段落计数，表内段落不计。黄色标注新增或重写；删除、移动、格式变化由以下映射记录，不是 Word 原生修订。')
    for c in changes:
        d.add_paragraph(c['id']+' '+c['type']+' '+c['old_location']+' → '+c['new_location'],'Heading 2')
        line('原文：'+c['before']);line('改文：'+(c['after'] or '〔删除〕'))
        if c['type']=='format':line('格式变更：'+json.dumps(c['format_requested'],ensure_ascii=False)+'；修改前后的完整属性保存在 change-map.json。')
        line('目的：'+c['purpose']+'；依据：'+c['evidence']+'；作者确认：'+('需要' if c['author_confirmation_needed'] else '本地核验，整体仍待作者审阅'))
    section('保留的边界与作者事项')
    for x in report['limitations']+report['author_questions']:line(x)
    section('实际检查与未验证范围');line('自动文件检查：'+checks['status']+'；'+checks['scope'])
    for x in report['checks']:line(x['name']+'：'+x['status']+'；'+x['detail'])
    line('本报告生成时视觉检查尚未执行。最终逐页结果见随交付的 visual-review.json；没有该文件表示未验收。作者最终认可未取得，未执行投稿。')
    d.save(path)
def build(source,manifest,out,clean=False):
    source=Path(source).resolve();out=Path(out).resolve();plan=json.loads(Path(manifest).read_text(encoding='utf-8-sig'))
    if out.exists():raise ValueError('Output directory must be new; do not overwrite a candidate')
    if sha(source.read_bytes())!=plan['input_sha256']:raise ValueError('Baseline SHA-256 differs')
    if plan.get('scope') not in ('full','local','continue'):raise ValueError('Build requires an editing scope')
    if not plan.get('operations'):raise ValueError('No edit operations; use inspect for diagnosis-only')
    infos,parts,root=read_package(source);revised,changes=apply_plan(root,plan,True)
    cleanroot,_=apply_plan(root,plan,False) if clean else (None,None)
    out.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='pef-build-',dir=out.parent) as td:
        temp=Path(td);review=temp/(source.stem+'_包装审阅版.docx')
        write_package(review,infos,parts,revised);checks=verify_pair(source,review)
        report=temp/(source.stem+'_诊断报告.docx');make_report(report,source,plan,changes,checks)
        if clean:write_package(temp/(source.stem+'_清洁候选稿.docx'),infos,parts,cleanroot)
        payload={'source':str(source),'source_sha256':plan['input_sha256'],'changes':changes,
          'files':{p.name:sha(p.read_bytes()) for p in temp.glob('*.docx')},'checks':checks,
          'visual_review':'NOT_RUN','author_acceptance':False,'publication':'NOT_PERFORMED'}
        (temp/'change-map.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
        # Recheck before publishing the independent local candidate directory.
        if sha(source.read_bytes())!=plan['input_sha256']:raise ValueError('Source changed during build')
        out.mkdir()
        for p in temp.iterdir():p.replace(out/p.name)
    return payload
def main():
    p=argparse.ArgumentParser(description=__doc__);sp=p.add_subparsers(dest='cmd',required=True)
    a=sp.add_parser('inspect');a.add_argument('source');a.add_argument('--out')
    a=sp.add_parser('build');a.add_argument('source');a.add_argument('manifest');a.add_argument('output_dir');a.add_argument('--clean',action='store_true')
    a=sp.add_parser('verify');a.add_argument('source');a.add_argument('output')
    args=p.parse_args()
    try:
        if args.cmd=='inspect':result=inventory(args.source)
        elif args.cmd=='build':result=build(args.source,args.manifest,args.output_dir,args.clean)
        else:result=verify_pair(args.source,args.output)
        output=json.dumps(result,ensure_ascii=False,indent=2)
        if getattr(args,'out',None):Path(args.out).write_text(output,encoding='utf-8')
        else:print(output)
    except (ValueError,KeyError,zipfile.BadZipFile) as exc:p.exit(2,'Refused: '+str(exc)+'\n')
if __name__=='__main__':main()
