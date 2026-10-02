"""Real DOCX media replacement: complete PNG/SVG binding and byte preservation."""
import copy
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as E

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import apply_figure_edits as F


def png(color):
    out = io.BytesIO()
    Image.new('RGB', (160, 80), color).save(out, format='PNG')
    return out.getvalue()


def svg(color):
    return ('<svg xmlns="http://www.w3.org/2000/svg" width="160" height="80" '
            'viewBox="0 0 160 80"><rect width="160" height="80" fill="' + color + '"/></svg>').encode()


class FigureEdits(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.dir = Path(self.temp.name)
        self.source = self.dir / 'source.docx'
        self.manifest = self.dir / 'manifest.json'
        self.dest = self.dir / 'candidate.docx'
        self.receipt = self.dir / 'receipt.json'
        self.old_png, self.old_svg = png('red'), svg('red')
        self.new_png, self.new_svg = png('blue'), svg('blue')
        (self.dir / 'new.png').write_bytes(self.new_png)
        (self.dir / 'new.svg').write_bytes(self.new_svg)
        self.fixture()

    def tearDown(self):
        self.temp.cleanup()

    def fixture(self, shared=False, header=False, alternative=True):
        doc = Document()
        paragraph = doc.add_paragraph('Keep all text, captions, fields and equations. ')
        paragraph.add_run('Styled variable').italic = True
        equation = OxmlElement('m:oMath')
        run = OxmlElement('m:r')
        value = OxmlElement('m:t')
        value.text = 'x + y'
        run.append(value)
        equation.append(run)
        paragraph._p.append(equation)
        for kind in ('begin', 'separate', 'end'):
            field_run = OxmlElement('w:r')
            field = OxmlElement('w:fldChar')
            field.set('{' + F.W + '}fldCharType', kind)
            field_run.append(field)
            paragraph._p.append(field_run)
        doc.add_picture(io.BytesIO(self.old_png), width=Inches(2), height=Inches(1))
        doc.add_paragraph('Figure 1. Preserved exact caption.')
        if shared:
            doc.add_picture(io.BytesIO(self.old_png), width=Inches(2), height=Inches(1))
        if header:
            doc.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(self.old_png), width=Inches(2))
        doc.save(self.source)
        with zipfile.ZipFile(self.source) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        root = E.fromstring(parts[F.DOCUMENT])
        if alternative:
            blip = root.find('.//a:blip', F.NS)
            ext_list = E.SubElement(blip, '{' + F.A + '}extLst')
            ext = E.SubElement(ext_list, '{' + F.A + '}ext', {'uri': '{96DAC541-7B7A-43D3-8B79-37D633B846F1}'})
            E.SubElement(ext, '{' + F.SVG + '}svgBlip', {'{' + F.R + '}embed': 'rIdSvgTest'})
            rels = E.fromstring(parts[F.DOCUMENT_RELS])
            E.SubElement(rels, '{' + F.REL + '}Relationship',
                         {'Id': 'rIdSvgTest', 'Type': F.IMAGE_REL, 'Target': 'media/image2.svg'})
            parts[F.DOCUMENT_RELS] = E.tostring(rels)
            types = E.fromstring(parts['[Content_Types].xml'])
            E.SubElement(types, '{' + F.CT + '}Default', {'Extension': 'svg', 'ContentType': 'image/svg+xml'})
            parts['[Content_Types].xml'] = E.tostring(types)
            parts['word/media/image2.svg'] = self.old_svg
        parts[F.DOCUMENT] = E.tostring(root)
        self.write_parts(parts)
        representations = [{'kind': 'primary', 'part': 'word/media/image1.png',
                            'old_media_sha256': F.sha(self.old_png), 'source_asset': 'new.png',
                            'source_asset_sha256': F.sha(self.new_png)}]
        if alternative:
            representations.append({'kind': 'svg', 'part': 'word/media/image2.svg',
                                    'old_media_sha256': F.sha(self.old_svg), 'source_asset': 'new.svg',
                                    'source_asset_sha256': F.sha(self.new_svg)})
        self.spec = {'source_sha256': F.sha(self.source.read_bytes()),
                     'replacements': [{'drawing_index': 1, 'representations': representations}]}

    def write_parts(self, parts):
        with zipfile.ZipFile(self.source, 'w', compression=zipfile.ZIP_DEFLATED) as z:
            z.comment = b'fixture package comment'
            for name, data in parts.items():
                z.writestr(name, data)

    def mutate_part(self, name, callback):
        with zipfile.ZipFile(self.source) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        parts[name] = callback(parts[name])
        self.write_parts(parts)
        self.spec['source_sha256'] = F.sha(self.source.read_bytes())

    def save_manifest(self):
        self.manifest.write_text(json.dumps(self.spec), encoding='utf-8')

    def run_apply(self):
        self.save_manifest()
        return F.apply(self.source, self.manifest, self.dest, self.receipt)

    def refused(self, pattern):
        original = self.source.read_bytes()
        with self.assertRaisesRegex((ValueError, OSError), pattern):
            self.run_apply()
        self.assertEqual(original, self.source.read_bytes())
        self.assertFalse(self.dest.exists())
        self.assertFalse(self.receipt.exists())

    def test_png_svg_update_preserves_every_other_member_and_protected_xml(self):
        original = self.source.read_bytes()
        result = self.run_apply()
        self.assertEqual(self.source.read_bytes(), original)
        with zipfile.ZipFile(self.source) as before, zipfile.ZipFile(self.dest) as after:
            self.assertEqual(before.namelist(), after.namelist())
            self.assertEqual(before.comment, after.comment)
            changed = []
            for a, b in zip(before.infolist(), after.infolist()):
                for key in ('date_time', 'compress_type', 'comment', 'extra', 'external_attr', 'internal_attr', 'create_system'):
                    self.assertEqual(getattr(a, key), getattr(b, key), (a.filename, key))
                if before.read(a.filename) != after.read(b.filename):
                    changed.append(a.filename)
            self.assertEqual(sorted(changed), ['word/media/image1.png', 'word/media/image2.svg'])
            self.assertEqual(after.read('word/media/image1.png'), self.new_png)
            self.assertEqual(after.read('word/media/image2.svg'), self.new_svg)
            self.assertEqual(before.read(F.DOCUMENT), after.read(F.DOCUMENT))
        self.assertIn('Preserved exact caption', Document(self.dest).paragraphs[-1].text)
        self.assertEqual(result['technical_status'], 'PASS')
        self.assertEqual(result['overall_status'], 'REVIEW_REQUIRED')
        self.assertEqual(result['quality_review'], 'NOT_PERFORMED')
        self.assertEqual(result['author_acceptance'], 'NOT_RECORDED')
        self.assertEqual(json.loads(self.receipt.read_text(encoding='utf-8')), result)

    def test_primary_only_is_supported_when_no_alternate_exists(self):
        self.fixture(alternative=False)
        self.assertEqual(self.run_apply()['changed_parts'], ['word/media/image1.png'])

    def test_missing_svg_refused_before_output(self):
        self.spec['replacements'][0]['representations'].pop()
        self.refused('representation')

    def test_source_hash_is_mandatory(self):
        self.spec.pop('source_sha256')
        self.refused('source_sha256')

    def test_stale_source_refused(self):
        self.spec['source_sha256'] = '0' * 64
        self.refused('Source DOCX SHA-256 mismatch')

    def test_old_media_hash_is_mandatory_and_checked(self):
        self.spec['replacements'][0]['representations'][0]['old_media_sha256'] = '0' * 64
        self.refused('Old media SHA-256 mismatch')
        self.spec['replacements'][0]['representations'][0].pop('old_media_sha256')
        self.refused('exact identity')

    def test_new_asset_hash_is_mandatory(self):
        self.spec['replacements'][0]['representations'][1].pop('source_asset_sha256')
        self.refused('exact identity')

    def test_stale_svg_asset_refused_before_png_write(self):
        (self.dir / 'new.svg').write_bytes(svg('green'))
        self.refused('New asset SHA-256 mismatch')

    def test_missing_svg_file_refused(self):
        (self.dir / 'new.svg').unlink()
        self.refused('new.svg')

    def test_destination_overwrite_refused(self):
        self.dest.write_bytes(b'keep destination')
        with self.assertRaisesRegex(ValueError, 'new files'):
            self.run_apply()
        self.assertEqual(self.dest.read_bytes(), b'keep destination')
        self.assertFalse(self.receipt.exists())

    def test_source_overwrite_refused(self):
        self.dest = self.source
        original = self.source.read_bytes()
        with self.assertRaisesRegex(ValueError, 'overwrite input'):
            self.run_apply()
        self.assertEqual(original, self.source.read_bytes())
        self.assertFalse(self.receipt.exists())

    def test_receipt_overwrite_and_output_collision_refused(self):
        self.receipt.write_bytes(b'keep receipt')
        with self.assertRaisesRegex(ValueError, 'new files'):
            self.run_apply()
        self.assertEqual(self.receipt.read_bytes(), b'keep receipt')
        self.assertFalse(self.dest.exists())
        self.receipt = self.dest
        self.refused('another output')

    def test_shared_unselected_drawing_refused(self):
        self.fixture(shared=True)
        self.refused('Shared media.*unselected')

    def test_shared_header_media_refused(self):
        self.fixture(header=True)
        self.refused('Shared media.*elsewhere')

    def test_duplicate_drawing_and_representation_refused(self):
        self.spec['replacements'].append(copy.deepcopy(self.spec['replacements'][0]))
        self.refused('repeated drawing_index')
        self.spec['replacements'].pop()
        reps = self.spec['replacements'][0]['representations']
        reps.append(copy.deepcopy(reps[0]))
        self.refused('duplicate representation')

    def test_selected_shared_media_still_refuses_repeated_replacement(self):
        self.fixture(shared=True)
        row = copy.deepcopy(self.spec['replacements'][0])
        row['drawing_index'] = 2
        row['representations'].pop()
        self.spec['replacements'].append(row)
        self.refused('Repeated media replacement')

    def test_wrong_media_part_identity_refused(self):
        self.spec['replacements'][0]['representations'][0]['part'] = 'word/media/another.png'
        self.refused('Media part identity')

    def test_extra_embedded_representation_refused(self):
        def change(data):
            root = E.fromstring(data)
            drawing = root.find('.//w:drawing', F.NS)
            E.SubElement(drawing, '{urn:test}unknownImage', {'{' + F.R + '}embed': 'rIdSvgTest'})
            return E.tostring(root)
        self.mutate_part(F.DOCUMENT, change)
        self.refused('Unsupported drawing relationship or representation')

    def test_unused_relationship_alias_to_selected_media_refused(self):
        def change(data):
            root = E.fromstring(data)
            E.SubElement(root, '{' + F.REL + '}Relationship',
                         {'Id': 'rIdAlias', 'Type': F.IMAGE_REL, 'Target': 'media/image1.png'})
            return E.tostring(root)
        self.mutate_part(F.DOCUMENT_RELS, change)
        self.refused('Shared media')

    def test_package_vml_part_and_signature_refused(self):
        for name, payload, expected in [('word/drawings/image.vml', b'<v:shape xmlns:v="urn:schemas-microsoft-com:vml"/>', 'VML'),
                                        ('_xmlsignatures/sig1.xml', b'<signature/>', 'Digitally signed')]:
            with self.subTest(name=name):
                self.fixture()
                with zipfile.ZipFile(self.source, 'a') as z:
                    z.writestr(name, payload)
                self.spec['source_sha256'] = F.sha(self.source.read_bytes())
                self.refused(expected)

    def test_external_and_escaping_relationships_refused(self):
        for target, mode, expected in [('https://example.org/a.svg', 'External', 'external'),
                                       ('../../outside.svg', 'Internal', 'escapes package')]:
            with self.subTest(target=target):
                self.fixture()
                def change(data):
                    tree = E.fromstring(data)
                    node = next(n for n in tree if n.get('Id') == 'rIdSvgTest')
                    node.set('Target', target)
                    node.set('TargetMode', mode)
                    return E.tostring(tree)
                self.mutate_part(F.DOCUMENT_RELS, change)
                self.refused(expected)

    def test_vml_and_alternatecontent_refused(self):
        for tag in ['{urn:schemas-microsoft-com:vml}shape',
                    '{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent']:
            with self.subTest(tag=tag):
                self.fixture()
                def change(data):
                    root = E.fromstring(data)
                    E.SubElement(root.find('w:body', F.NS), tag)
                    return E.tostring(root)
                self.mutate_part(F.DOCUMENT, change)
                self.refused('VML or AlternateContent')

    def test_png_extension_does_not_allow_non_png_bytes(self):
        (self.dir / 'new.png').write_bytes(self.new_svg)
        self.spec['replacements'][0]['representations'][0]['source_asset_sha256'] = F.sha(self.new_svg)
        self.refused('cannot identify image')

    def test_extension_and_content_type_incompatibility_refused(self):
        self.spec['replacements'][0]['representations'][0]['source_asset'] = 'new.svg'
        self.refused('extension incompatibility')
        self.fixture()
        self.mutate_part('[Content_Types].xml', lambda b: b.replace(b'image/png', b'image/jpeg'))
        self.refused('content type incompatibility')

    def test_unchanged_svg_is_not_a_completed_replacement(self):
        (self.dir / 'new.svg').write_bytes(self.old_svg)
        self.spec['replacements'][0]['representations'][1]['source_asset_sha256'] = F.sha(self.old_svg)
        self.refused('still contains old media')

    def test_asset_ratio_must_fit_existing_extent(self):
        payload = self.new_svg.replace(b'width="160"', b'width="80"')
        (self.dir / 'new.svg').write_bytes(payload)
        self.spec['replacements'][0]['representations'][1]['source_asset_sha256'] = F.sha(payload)
        self.refused('aspect ratio')

    def test_cli_help_and_actual_complete_replacement(self):
        script = ROOT / 'scripts' / 'apply_figure_edits.py'
        help_result = subprocess.run([sys.executable, str(script), '--help'], capture_output=True, text=True)
        self.assertEqual(help_result.returncode, 0)
        self.assertIn('old_media_sha256', help_result.stdout)
        self.save_manifest()
        run = subprocess.run([sys.executable, str(script), str(self.source), str(self.manifest),
                              str(self.dest), '--receipt', str(self.receipt)], capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertEqual(json.loads(run.stdout)['technical_status'], 'PASS')
        self.assertTrue(self.dest.exists())


if __name__ == '__main__':
    unittest.main()
