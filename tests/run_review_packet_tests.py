"""Test exact review selection, source binding and leakage refusal."""
import argparse,copy,json,sys
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from make_review_packet import build,sha
from review_docx import inventory

def main():
    p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);a=p.parse_args();a.out.mkdir(parents=True,exist_ok=False)
    src=a.out/'input.docx';d=Document();d.add_paragraph('Heading');d.add_paragraph('Selected claim.');d.add_paragraph('SECRET_FULLTEXT_SENTINEL');d.save(src)
    inv=inventory(src);sel={'source_sha256':inv['sha256'],'purpose':'SECRET_EDIT_RATIONALE','paragraphs':[{'id':x['id'],'text_sha256':x['text_sha256']} for x in inv['paragraphs'][:2]]}
    results=[]
    def check(name,ok):results.append({'name':name,'status':'PASS' if ok else 'FAIL'})
    r=build(src,sel,a.out/'packet');s=(a.out/'packet/reading.md').read_text(encoding='utf-8')
    check('selected prose copied',all(t in s for t in ('Heading','Selected claim.')))
    check('excluded prose absent from every output',all(b'SECRET_FULLTEXT_SENTINEL' not in f.read_bytes() for f in (a.out/'packet').iterdir()))
    check('editor purpose excluded from reviewer files',all(b'SECRET_EDIT_RATIONALE' not in f.read_bytes() for f in (a.out/'packet').iterdir()))
    check('full source not copied',not list((a.out/'packet').glob('*.docx')))
    check('source unchanged',sha(src.read_bytes())==inv['sha256'])
    check('tool does not claim blind review',r['blind_review_status']=='NOT_ESTABLISHED' and r['reading_status']=='NOT_READ')
    for name,mut in [('wrong source',lambda v:v.update(source_sha256='0'*64)),('wrong paragraph',lambda v:v['paragraphs'][0].update(text_sha256='0'*64)),('unknown paragraph',lambda v:v['paragraphs'][0].update(id='P9999')),('duplicate paragraph',lambda v:v['paragraphs'].append(v['paragraphs'][0])),('empty scope',lambda v:v.update(paragraphs=[])),('unknown selector',lambda v:v.update(stop_at='Introduction'))]:
        v=copy.deepcopy(sel);mut(v);dest=a.out/name.replace(' ','-')
        try:build(src,v,dest);ok=False
        except ValueError:ok=not dest.exists()
        check(name+' refused before output',ok)
    try:build(src,sel,a.out/'packet');ok=False
    except ValueError:ok=True
    check('existing output protected',ok)
    result={'status':'PASS' if all(x['status']=='PASS' for x in results) else 'FAIL','count':len(results),'tests':results}
    (a.out/'results.json').write_text(json.dumps(result,indent=2),encoding='utf-8');print(json.dumps(result));return result['status']!='PASS'
if __name__=='__main__':raise SystemExit(main())
