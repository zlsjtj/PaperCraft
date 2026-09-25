"""Deterministic teaching-fixture tests. No real measurements or manuscript edits."""
import argparse,copy,importlib.util,json,sys,zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from lxml import etree as E
from PIL import Image,ImageDraw,ImageFont

ROOT=Path(__file__).resolve().parents[1]
def load(name,path):
    s=importlib.util.spec_from_file_location(name,path);m=importlib.util.module_from_spec(s);s.loader.exec_module(m);return m
H=load('review_docx',ROOT/'scripts/review_docx.py')
DEP=load('check_dependencies',ROOT/'scripts/check_dependencies.py')

def make_fixture(work):
    work.mkdir(parents=True,exist_ok=False)
    img=Image.new('RGB',(960,160),'white');draw=ImageDraw.Draw(img)
    font=ImageFont.load_default(size=28)
    for x,label in [(20,'BUILD PLAN'),(350,'REUSE PLAN'),(680,'RUN BATCH')]:
        draw.rounded_rectangle((x,30,x+245,130),radius=8,fill='#e8eef5',outline='#405b76',width=2)
        draw.text((x+18,64),label,fill='#192a3c',font=font)
    draw.line((265,80,340,80),fill='#405b76',width=3);draw.line((595,80,670,80),fill='#405b76',width=3)
    for x in (340,670):draw.polygon([(x,80),(x-12,73),(x-12,87)],fill='#405b76')
    img.save(work/'teaching-diagram.png')
    d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.7)
    d.styles['Normal'].font.name='Times New Roman';d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.space_after=Pt(6)
    for name,size in [('Title',20),('Heading 1',12)]:
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(6)
        for el in st._element.xpath('.//w:pBdr'):el.getparent().remove(el)
    texts=[('Title','An efficient execution method'),('Normal','TEACHING FIXTURE ONLY. All timings below are invented for testing document behavior, not research evidence.'),
      ('Heading 1','Abstract'),('Normal','We present an efficient method and evaluate it. The method improves performance.'),
      ('Heading 1','Introduction'),('Normal','Our method has several optimizations. It uses an execution plan.'),
      ('Heading 1','Methods'),('Normal','The batch size is 32. Its selection record and boundary checks are unavailable.'),
      ('Normal','The teaching design builds a lookup plan once and reuses it for each batch of unchanged-shape inputs. Changed shapes require a new plan.')]
    for style,t in texts:d.add_paragraph(t,style)
    d.paragraphs[1].runs[0].font.highlight_color=WD_COLOR_INDEX.BRIGHT_GREEN
    p=d.add_paragraph();math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='s = T_control / T_candidate';r.append(t);math.append(r);p._p.append(math)
    d.add_picture(str(work/'teaching-diagram.png'),width=Inches(5.5))
    d.add_paragraph('Figure 1. Schematic of the supplied teaching design; arrows denote processing order.')
    d.add_paragraph('Results','Heading 1')
    d.add_paragraph('The three invented workloads A, B and C each have five repeated timings. Median control/candidate milliseconds are 10/8, 10/10 and 10/12. There are three workloads, not fifteen independent samples.')
    table=d.add_table(rows=1,cols=3);table.style='Table Grid'
    for c,t in zip(table.rows[0].cells,['Workload','Control ms','Candidate ms']):c.text=t
    for row in [('A','10','8'),('B','10','10'),('C','10','12')]:
        for c,t in zip(table.add_row().cells,row):c.text=t
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:p.paragraph_format.space_after=Pt(0)
    d.add_paragraph('Discussion','Heading 1');d.add_paragraph('Some cases lose. More work is needed.')
    d.add_paragraph('Our method is an efficient method with several optimizations.')
    d.add_paragraph('Conclusion','Heading 1');d.add_paragraph('The method is useful and efficient.')
    d.add_paragraph('The cause of the loss on C has not been measured. No cache counters or additional experiments are available.')
    d.add_paragraph('Declarations','Heading 1');d.add_paragraph('Author information, funding and public-data status remain pending. No author approval is asserted.')
    p=d.add_paragraph('Source locator ');b=OxmlElement('w:bookmarkStart');b.set(H.Q('id'),'7');b.set(H.Q('name'),'TeachingSource');p._p.append(b)
    p.add_run('retained');b=OxmlElement('w:bookmarkEnd');b.set(H.Q('id'),'7');p._p.append(b)
    p=d.add_paragraph('Cross reference: ');f=OxmlElement('w:fldSimple');f.set(H.Q('instr'),' REF TeachingSource ')
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='retained';r.append(t);f.append(r);p._p.append(f)
    p=d.add_paragraph('Protected footnote');r=OxmlElement('w:r');n=OxmlElement('w:footnoteReference');n.set(H.Q('id'),'1');r.append(n);p._p.append(r)
    out=work/'teaching_manuscript.docx';d.save(out)
    infos,parts,_=H.read_package(out)
    foot=E.Element(H.Q('footnotes'),nsmap={'w':H.W})
    for ident,typ in [(-1,'separator'),(0,'continuationSeparator'),(1,None)]:
        n=E.SubElement(foot,H.Q('footnote'));n.set(H.Q('id'),str(ident))
        if typ:n.set(H.Q('type'),typ)
        p=E.SubElement(n,H.Q('p'));r=E.SubElement(p,H.Q('r'))
        if typ:E.SubElement(r,H.Q(typ))
        else:E.SubElement(r,H.Q('t')).text='Teaching note; preserve this exact source text.'
    parts['word/footnotes.xml']=E.tostring(foot,xml_declaration=True,encoding='UTF-8')
    rels=E.fromstring(parts['word/_rels/document.xml.rels']);r=E.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship')
    r.set('Id','rIdPEFFootnote');r.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes');r.set('Target','footnotes.xml')
    parts['word/_rels/document.xml.rels']=E.tostring(rels)
    ct=E.fromstring(parts['[Content_Types].xml']);c=E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override')
    c.set('PartName','/word/footnotes.xml');c.set('ContentType','application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    parts['[Content_Types].xml']=E.tostring(ct)
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for name,content in parts.items():z.writestr(name,content)
    return out

def manifest(source):
    inv=H.inventory(source);paras=inv['paragraphs']
    def get(prefix):return next(x for x in paras if x['text'].startswith(prefix))
    ops=[]
    def op(kind,prefix,**kw):
        p=get(prefix);ops.append({'id':f'E{len(ops)+1:03d}','kind':kind,'target':p['id'],'expect':p['text'],
          'purpose':kw.pop('purpose','把教学材料已经提供的信息写清楚'),'evidence':kw.pop('evidence','教学原稿 Methods/Results；不是科研实测'),'confirm':kw.pop('confirm',False),**kw})
    op('replace','An efficient',text='Reusing execution plans across unchanged-shape batches')
    op('replace','We present',text='The teaching design reuses one lookup plan across unchanged-shape batches, avoiding repeated plan construction. Invented test workloads show a benefit on A, parity on B and a loss on C. The cause of the loss and the batch-size selection remain unverified.')
    op('replace','Our method has',text='Repeated setup can duplicate work when successive inputs have the same shape. The supplied design retains a lookup plan across these batches and rebuilds it when the shape changes. Its contribution here is this conditional reuse rule; priority over prior research has not been established.')
    op('insert_after','The teaching design',text='Plan reuse requires the shape to remain unchanged; this condition defines when the stored mapping can be used. The supplied material does not establish optimal batch size or completed boundary validation.',confirm=True)
    op('format','Figure 1.',format={'space_after_pt':8,'keep_next':False},purpose='改善图注与后文间距')
    op('replace','Some cases lose.',text='The three teaching cases separate an opportunity from its limits: reusing setup may help A, but the loss on C prevents a uniform performance claim. The current records do not isolate setup, memory traffic or computation costs. Choosing between reuse and reconstruction requires further evidence.',confirm=True)
    op('delete','Our method is an',purpose='删除重复宣传句；原文完整保留在变化映射')
    op('replace','The method is useful',text='Reusing a lookup plan offers a concrete execution option for unchanged-shape batches. The teaching outcomes delimit its use rather than establish universal superiority. Parameter provenance and the cause of the adverse case remain open.')
    op('move_after','The cause of the loss',after=get('Some cases lose.')['id'],purpose='将已有未解原因移到 Discussion，保持事实不变')
    report=json.loads((ROOT/'templates/revision-manifest.json').read_text(encoding='utf-8'))['report']
    report.update(summary='教学样例验收。已实际改写标题、摘要、引言、讨论和结论，补充条件解释，保留负结果；全部时间数字为虚构测试数据，不能用于论文。',
      thesis='将重复构造 lookup plan 的工作改为相同形状批次之间复用，条件是形状不变；教学结果包含收益、中性和损失。',
      titles=[{'title':'Reusing execution plans across unchanged-shape batches','reason':'推荐，突出可核查机制与条件'}, {'title':'Conditional plan reuse for batched execution','reason':'强调适用条件而非首创或最佳性能'}],
      dimensions=[{'name':'创新','before':'只有 efficient 等泛称','after':'明确复用对象与形状条件'}, {'name':'工作量','before':'参数和验证不清，讨论过短','after':'展开条件与成本问题；缺记录仍待确认'}, {'name':'面子工程','before':'重复宣传、未解原因位置分散','after':'删除重复、移动原段并调整图注间距'}],
      ledger=[{'id':'C01','claim':'同形状复用计划','difference':'教学设计的明确操作；最近邻文献未提供','source':'教学 Methods 段','nature':'原文报告','status':'完成','scope':'仅教学样例，原理未作实测验证'}, {'id':'C02','claim':'结果有利与不利并存','difference':'相同三个教学工作负载比较','source':'Results 段与表 1','nature':'原文报告','status':'完成','scope':'虚构数据，仅验证保留口径'}],
      workload=[{'name':'背景','detail':'重复构造相同形状计划的开销'}, {'name':'研究现状','detail':'未提供文献，不宣称首创'}, {'name':'参数','detail':'32 来源未知，不编标定'}, {'name':'实验或模型','detail':'表中三案例是教学数据，边界检查缺失'}, {'name':'重复','detail':'每工作负载五次，共三独立工作负载，不写成十五'}, {'name':'验证','detail':'实际执行文档对象保留与高亮检查；未执行科研实验。新增为文字解释，没有新增测量'}],
      figures=[{'label':'图 1','status':'保留','detail':'保留给定教学流程图及源图；调整图注间距。图内容正确性与排版另做视觉核查'}],
      limitations=['A/B/C 三种结果全部保留；C 的原因未测，不能归因 cache。','参数来源、边界验证及最近邻文献尚缺。'],
      author_questions=['确认参数 32 的原始选择记录和实际边界测试；不得把此教学稿当真实研究稿。'],
      checks=[{'name':'数据','status':'教学材料一致性','detail':'数字未改；并非实际 benchmark 复核'}, {'name':'引用','status':'未执行','detail':'教学样例未给外部文献，不制造引用'}])
    return {'input_sha256':inv['sha256'],'scope':'full','execution':'sequential','operations':ops,'report':report}

def run(work):
    source=make_fixture(work);plan=manifest(source);(work/'manifest.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    results=[]
    def case(name,func):
        try:detail=func();results.append({'case':name,'status':'PASS','detail':detail})
        except Exception as e:results.append({'case':name,'status':'FAIL','detail':repr(e)})
    def refuses(mut,name):
        p=copy.deepcopy(plan);mut(p)
        try:H.build(source,write_plan(p,name),work/name)
        except ValueError as e:return str(e)
        raise AssertionError('Unsafe input not refused')
    def write_plan(p,name):
        f=work/(name+'.json');f.write_text(json.dumps(p,ensure_ascii=False),encoding='utf-8');return f
    initial=H.sha(source.read_bytes())
    case('dual_word_build',lambda:H.build(source,work/'manifest.json',work/'full-output',True)['checks'])
    out=work/'full-output';review=out/'teaching_manuscript_包装审阅版.docx';clean=out/'teaching_manuscript_清洁候选稿.docx'
    def preserved():
        r=H.verify_pair(source,review);assert r['objects']['m:oMath']==1 and r['objects']['w:tbl']==1 and r['objects']['w:drawing']==1
        assert r['objects']['w:footnoteReference']==1 and r['objects']['w:fldSimple']==1;return r
    case('equation_table_figure_footnote_field_bookmark_preservation',preserved)
    def highlights():
        _,_,rr=H.read_package(review);_,_,cr=H.read_package(clean)
        yellow=rr.xpath('//w:highlight[@w:val="yellow"]',namespaces=H.NS)
        assert len(yellow)==sum(o['kind'] in ('replace','insert_after') for o in plan['operations'])
        assert not cr.xpath('//w:highlight[@w:val="yellow"]',namespaces=H.NS)
        assert len(cr.xpath('//w:highlight[@w:val="green"]',namespaces=H.NS))==1
        return 'New prose yellow; clean candidate retains pre-existing green highlight'
    case('new_highlights_and_existing_highlight_preservation',highlights)
    def mappings():
        m=json.loads((out/'change-map.json').read_text(encoding='utf-8'));paras=H.inventory(review)['paragraphs'];by={p['id']:p['text'] for p in paras}
        assert len(m['changes'])==len(plan['operations'])
        for c in m['changes']:
            if c['type']=='delete':assert c['new_location']=='deleted' and c['before'] not in by.values()
            else:assert by[c['new_location']]==c['after']
        assert {'delete','move_after','format'}<=set(c['type'] for c in m['changes'])
        assert m['author_acceptance'] is False and m['visual_review']=='NOT_RUN'
        return 'All edit IDs and new positions resolve; deletion, movement and format explicitly mapped'
    case('complete_change_mapping_and_status_separation',mappings)
    def schema_order():
        _,_,rr=H.read_package(review)
        caption=next(p for p in H.paragraphs(rr) if H.text(p).startswith('Figure 1.'))
        tags=[E.QName(x).localname for x in caption.find('w:pPr',H.NS)]
        assert tags.index('keepNext')<tags.index('spacing')
        probe=E.fromstring(f'<w:p xmlns:w="{H.W}"><w:r><w:rPr><w:sz w:val="22"/><w:lang w:val="en-US"/></w:rPr><w:t>text</w:t></w:r></w:p>')
        H.write_text(probe,'new text',True)
        tags=[E.QName(x).localname for x in probe.find('w:r/w:rPr',H.NS)]
        assert tags.index('sz')<tags.index('highlight')<tags.index('lang')
        return 'Supported formatting properties use OOXML schema order'
    case('supported_OOXML_property_order',schema_order)
    case('wrong_baseline_refused',lambda:refuses(lambda p:p.update(input_sha256='0'*64),'bad-hash'))
    case('wrong_expected_text_refused',lambda:refuses(lambda p:p['operations'][0].update(expect='wrong'),'bad-text'))
    def reject_object(prefix,name):
        target=next(p for p in H.inventory(source)['paragraphs'] if p['text'].startswith(prefix))
        return refuses(lambda p:p.update(operations=[dict(p['operations'][0],target=target['id'],expect=target['text'])]),name)
    case('native_equation_edit_refused',lambda:reject_object('s =','bad-equation'))
    case('field_edit_refused',lambda:reject_object('Cross reference:','bad-field'))
    case('existing_highlight_edit_refused',lambda:reject_object('TEACHING FIXTURE','bad-highlight'))
    def structural_refusal(mode):
        _,_,root=H.read_package(source);p=H.paragraphs(root)[0]
        if mode=='mixed':
            r=E.SubElement(p,H.Q('r'));rp=E.SubElement(r,H.Q('rPr'));E.SubElement(rp,H.Q('b'));E.SubElement(r,H.Q('t')).text=' bold'
        else:E.SubElement(p,H.Q('ins')).set(H.Q('id'),'9')
        q=copy.deepcopy(plan);q['operations']=[dict(q['operations'][0],expect=H.text(p))]
        try:H.apply_plan(root,q)
        except ValueError as e:return str(e)
        raise AssertionError('Complex structure not refused')
    case('mixed_run_formatting_refused',lambda:structural_refusal('mixed'))
    case('existing_native_revisions_refused',lambda:structural_refusal('tracked'))
    def no_overwrite():
        try:H.build(source,work/'manifest.json',out)
        except ValueError:return 'Existing candidate protected'
        raise AssertionError('Overwrote existing candidate')
    case('candidate_overwrite_refused',no_overwrite)
    case('diagnosis_mode_mutation_refused',lambda:refuses(lambda p:p.update(scope='diagnosis'),'bad-mode'))
    def local():
        p=copy.deepcopy(plan);p['scope']='local';p['operations']=[p['operations'][2]]
        H.build(source,write_plan(p,'local'),work/'local-output')
        _,a,ar=H.read_package(source);_,b,br=H.read_package(work/'local-output/teaching_manuscript_包装审阅版.docx')
        changes=[i for i,(x,y) in enumerate(zip(H.paragraphs(ar),H.paragraphs(br))) if H.canonical(x)!=H.canonical(y)]
        assert changes==[5];return 'Exactly requested paragraph P0006 changed; other paragraphs retained'
    case('local_scope_changes_one_paragraph',local)
    def continuation():
        inv=H.inventory(clean);p=copy.deepcopy(plan);p['input_sha256']=inv['sha256'];p['scope']='continue'
        p['operations']=[dict(p['operations'][0],expect=inv['paragraphs'][0]['text'],text='Conditional reuse of execution plans across unchanged-shape batches')]
        H.build(clean,write_plan(p,'continue'),work/'continue-output')
        c=H.inventory(work/'continue-output/teaching_manuscript_清洁候选稿_包装审阅版.docx')
        assert [x['text'] for x in c['paragraphs'][1:]]==[x['text'] for x in inv['paragraphs'][1:]]
        return 'New parent hash and inventory; accepted prior content retained'
    case('continuation_preserves_prior_edits',continuation)
    def missing():
        r=work/'missing-skill';r.mkdir();(r/'SKILL.md').write_text('# Entry\n[Required](references/missing.md)',encoding='utf-8')
        result=DEP.check(r);assert result['missing_local_links'];return result['missing_local_links']
    case('missing_reference_detected',missing)
    case('source_immutable',lambda:'Original SHA-256 unchanged' if H.sha(source.read_bytes())==initial else (_ for _ in ()).throw(AssertionError('Source changed')))
    (work/'test-results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'passed':sum(x['status']=='PASS' for x in results),'total':len(results),'results':results},ensure_ascii=False,indent=2))
    return all(x['status']=='PASS' for x in results)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--work-dir',type=Path,required=True);a=p.parse_args()
    raise SystemExit(0 if run(a.work_dir.resolve()) else 1)
