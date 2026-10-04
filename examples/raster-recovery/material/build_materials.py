"""Rebuild the synthetic input package. Does not read or write evaluator files."""
from __future__ import annotations
import csv
import hashlib
import io
import json
import math
import re
import sys
import argparse
import shutil
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

sys.dont_write_bytecode = True
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from model import (PARAMETERS, TRACES, anchor_control, run_tests, simulate)

ROOT = Path(__file__).resolve().parent
FONT = None
BOLD_FONT = None
FIXED_DATE = datetime(2026, 10, 2, 0, 0, 0)

def write_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

def font(size, bold=False):
    return ImageFont.truetype(str(BOLD_FONT if bold else FONT), size)

def arrow(draw, a, b, fill="#737373", width=2, head=10):
    draw.line([a, b], fill=fill, width=width)
    angle = math.atan2(b[1] - a[1], b[0] - a[0])
    p1 = (b[0] - head * math.cos(angle - .45), b[1] - head * math.sin(angle - .45))
    p2 = (b[0] - head * math.cos(angle + .45), b[1] - head * math.sin(angle + .45))
    draw.polygon([b, p1, p2], fill=fill)

def draw_raster():
    im = Image.new("RGB", (1500, 1050), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 18), "Original sketch 1  |  SYNTHETIC DEMO  |  sample centers and one reset state", font=font(29, True), fill="black")
    x0, y0, sx, sy = 170, 125, 76, 55
    for c in range(16):
        d.text((x0 + c*sx - 8, 76), str(c), font=font(23), fill="#333333")
    d.text((40, 76), "column", font=font(23), fill="black")
    for r in range(12):
        y = y0 + r*sy
        d.text((35, y-13), f"r={r:02}", font=font(24), fill="black")
        seq = list(range(16)) if r % 2 == 0 else list(reversed(range(16)))
        for i in range(15):
            c0, c1 = seq[i:i+2]
            sign = 1 if c1 > c0 else -1
            arrow(d, (x0+c0*sx+sign*20, y), (x0+c1*sx-sign*21, y), width=2, head=8)
        if r < 11:
            c = seq[-1]
            arrow(d, (x0+c*sx, y+18), (x0+c*sx, y+sy-19), width=2, head=8)
        for c in range(16):
            fill = "#dadada" if r < 5 else "#f4f4f4"
            if r == 5:
                if c >= 8:
                    fill = "#8fbeb7"
                elif c in (6, 7):
                    fill = "#f2d58c"
                elif c == 5:
                    fill = "#e99c9c"
            x = x0+c*sx
            d.rectangle((x-19, y-16, x+19, y+16), fill=fill, outline="#555555", width=1)
        if r == 5:
            for b, c_left in ((0, 12), (1, 8)):
                d.rectangle((x0+c_left*sx-27, y-23, x0+(c_left+3)*sx+27, y+23), outline="#356d66", width=3)
                d.text((x0+c_left*sx+65, y-43), f"b{b}", font=font(19), fill="#356d66")
            d.ellipse((x0+15*sx-25, y-25, x0+15*sx+25, y+25), outline="#111111", width=3)
            d.rectangle((x0+7*sx-25, y-25, x0+7*sx+25, y+25), outline="#705798", width=4)
    d.text((45, 785), "Grey rows are complete.  Green = sealed.  Yellow = volatile.  Red = interrupted acquisition.", font=font(25), fill="black")
    d.text((45, 827), "Row 5 runs right to left: sealed k=0..7; volatile k=8,9; interrupted k=10 (physical c=5).", font=font(25), fill="black")
    d.text((45, 869), "Restart q=88: r=5, k=8, physical c=7, x=28 mm, y=30 mm.  Entry reference is c=15, x=60 mm.", font=font(25), fill="black")
    d.text((45, 911), "Pitch: x=4 mm, y=6 mm.  Sample-center extent: 60 x 66 mm.  Blocks group traversal order.", font=font(25), fill="black")
    d.text((45, 972), "Constructed state only; not a timed capture from a result trace.  Sources S1 and S2.", font=font(22), fill="#555555")
    im.save(ROOT / "original_figures" / "figure_1_original.png", dpi=(238.125, 238.125))

def box(draw, bounds, title, lines):
    draw.rectangle(bounds, fill="#efefef", outline="#444444", width=3)
    x, y = bounds[:2]
    draw.text((x+15, y+12), title, font=font(27, True), fill="#222222")
    for i, line in enumerate(lines):
        draw.text((x+15, y+54+i*30), line, font=font(23), fill="#222222")

def draw_storage():
    im = Image.new("RGB", (1500, 1050), "white")
    d = ImageDraw.Draw(im)
    d.text((35, 20), "Original sketch 2  |  SYNTHETIC DEMO  |  BS forward work and restart", font=font(30, True), fill="black")
    d.text((40, 68), "Restart order: home (120 ms), recovery reader, row anchor, then acquisition.", font=font(23), fill="#444444")
    box(d, (40, 110, 405, 320), "Traversal state", ["q = 16r + k", "odd row: c = 15-k", "physical position + ordinal", "acquire: 10 ms each"])
    box(d, (535, 110, 940, 320), "Volatile buffer", ["four records, 1024 bytes", "cut discards this buffer", "k = 4b ... 4b+3", "block b = 0,1,2,3"])
    box(d, (1080, 110, 1460, 350), "Persistent blocks", ["payload then seal", "epoch / row / block", "count / length / CRC", "1048 bytes per block", "2 ms combined event"])
    box(d, (1075, 475, 1460, 710), "Row checkpoint", ["slot A       slot B", "64 bytes    64 bytes", "generation + checksum", "update after four blocks", "1 ms; inherited slots"])
    box(d, (515, 475, 950, 735), "Recovery reader", ["valid checkpoint -> row", "inspect blocks b=0 to 3", "match metadata + CRC", "stop at first invalid block", "accept contiguous prefix", "lookup: 6 ms after a cut"])
    box(d, (40, 475, 420, 735), "Position establishment", ["row anchor: 30 ms", "travel to first missing q", "known physical coordinate", "must precede acquisition", "after reset", "home already completed"])
    arrow(d, (405, 205), (535, 205), width=3)
    d.text((423, 161), "sample", font=font(20), fill="black")
    arrow(d, (940, 205), (1080, 205), width=3)
    d.text((957, 158), "write", font=font(20), fill="black")
    arrow(d, (1266, 350), (1266, 475), width=3)
    d.text((1288, 386), "row end", font=font(21), fill="black")
    arrow(d, (1075, 599), (950, 599), width=3)
    d.text((963, 553), "read", font=font(21), fill="black")
    arrow(d, (1120, 350), (909, 475), width=3)
    d.text((920, 382), "read seals", font=font(21), fill="black")
    arrow(d, (515, 600), (420, 600), width=3)
    d.text((426, 550), "first q", font=font(20), fill="black")
    arrow(d, (225, 475), (225, 320), width=3)
    d.text((45, 386), "resume", font=font(22), fill="black")
    d.text((40, 797), "Accepted output does not prove physical position. CRC-valid bytes may describe the wrong location.", font=font(25), fill="black")
    d.text((40, 845), "All four blocks valid but row cursor old: finish row checkpoint, then enter the next row.", font=font(25), fill="black")
    d.text((40, 893), "At a gap, later valid seals are ignored. Cut before seal publication leaves that block unaccepted.", font=font(25), fill="black")
    d.text((40, 956), "Diagram records assumed behavior, not a measured device. Sources S1, S2, and S4.", font=font(23), fill="#555555")
    im.save(ROOT / "original_figures" / "figure_2_original.png", dpi=(238.125, 238.125))

def math_run(text):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r

def math_sub(base, sub):
    obj = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    e.append(math_run(base))
    s = OxmlElement("m:sub")
    s.append(math_run(sub))
    obj.extend([e, s])
    return obj

def math_sum(subtext, expression):
    obj = OxmlElement("m:nary")
    pr = OxmlElement("m:naryPr")
    char = OxmlElement("m:chr")
    char.set(qn("m:val"), "∑")
    limit = OxmlElement("m:limLoc")
    limit.set(qn("m:val"), "subSup")
    sup_hide = OxmlElement("m:supHide")
    sup_hide.set(qn("m:val"), "1")
    pr.extend([char, limit, sup_hide])
    sub = OxmlElement("m:sub")
    sub.append(math_run("e∈"))
    sub.append(math_sub("E", subtext))
    sup = OxmlElement("m:sup")
    exp = OxmlElement("m:e")
    exp.append(math_run(expression))
    obj.extend([pr, sub, sup, exp])
    return obj

def add_equation(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_para = OxmlElement("m:oMathPara")
    obj = OxmlElement("m:oMath")
    obj.append(math_sub("T", "power"))
    obj.append(math_run(" = "))
    obj.append(math_sum("done", "d(e)"))
    obj.append(math_run(" + "))
    obj.append(math_sum("cut", "τ(e)"))
    obj.append(math_run("      (1)"))
    math_para.append(obj)
    p._p.append(math_para)

def add_table(doc, result_map):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [25, 25, 36, 37, 37]
    for c, text in zip(table.rows[0].cells, ["Trace", "Cuts", "RC (ms)", "AR (ms)", "BS (ms)"]):
        c.text = text
    for trace, cuts in TRACES.items():
        cells = table.add_row().cells
        values = [trace, str(len(cuts))] + [f"{result_map[trace, p]['powered_ms']:.1f}" for p in ("RC", "AR", "BS")]
        for c, value in zip(cells, values):
            c.text = value
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EAEAEA")
                cell._tc.get_or_add_tcPr().append(shade)
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + side)
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn("w:" + attr), value)
        borders.append(element)
    table._tbl.tblPr.append(borders)

def save_deterministic(doc, path):
    buf = io.BytesIO()
    doc.save(buf)
    with ZipFile(buf) as src, ZipFile(path, "w", ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            info = ZipInfo(name, date_time=(2026, 10, 2, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, src.read(name))

def build_docx(result_map):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(20)
    sec.left_margin = sec.right_margin = Mm(25)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(18)
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Caption"].font.size = Pt(9)
    for paragraph in (ROOT / "manuscript_source.md").read_text(encoding="utf-8").split("\n\n"):
        text = paragraph.strip()
        if not text:
            continue
        if text.startswith("# "):
            doc.add_paragraph(text[2:], "Title")
        elif text.startswith("## "):
            doc.add_heading(text[3:], level=1)
        elif text == "[[FIGURE_1]]" or text == "[[FIGURE_2]]":
            n = 1 if text == "[[FIGURE_1]]" else 2
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(ROOT / "original_figures" / f"figure_{n}_original.png"), width=Mm(160))
            for element in run._r.xpath(".//wp:docPr"):
                element.set("descr", "Original synthetic DEMO raster state sketch" if n == 1 else "Original synthetic DEMO block storage and recovery sketch")
        elif text == "[[TABLE_1]]":
            add_table(doc, result_map)
        elif text == "[[EQ1]]":
            add_equation(doc)
        elif text.startswith("Figure ") or text.startswith("Table "):
            doc.add_paragraph(text, "Caption")
        else:
            doc.add_paragraph(text)
    doc.core_properties.title = "Restart notes for a twelve row raster acquisition job"
    doc.core_properties.author = "Synthetic DEMO material author"
    doc.core_properties.subject = "Constructed input for a manuscript and scientific figure editing trial"
    doc.core_properties.created = FIXED_DATE
    doc.core_properties.modified = FIXED_DATE
    save_deterministic(doc, ROOT / "manuscript_input.docx")

def count_words(text):
    return len(re.findall(r"\b[A-Za-z0-9]+(?:['-][A-Za-z0-9]+)*\b", text))

def main():
    for sub in ("data", "tests", "original_figures"):
        (ROOT / sub).mkdir(exist_ok=True)
    results, all_events = [], []
    for trace, cuts in TRACES.items():
        for policy in ("RC", "AR", "BS"):
            result, events = simulate(policy, cuts, trace)
            results.append(result)
            all_events.extend(events)
    result_map = {(r["trace"], r["policy"]): r for r in results}
    write_json(ROOT / "data" / "parameters.json", PARAMETERS)
    write_json(ROOT / "data" / "traces.json", {"evidence_kind": "constructed deterministic schedules", "clock": "cumulative powered milliseconds", "traces": TRACES})
    write_csv(ROOT / "data" / "results.csv", results)
    write_csv(ROOT / "data" / "events.csv", all_events)
    write_csv(ROOT / "data" / "anchor_control.csv", anchor_control())
    write_csv(ROOT / "data" / "storage.csv", [
        {"policy": "RC", "payload_buffer_bytes": 4096, "persistent_bytes": 49280, "expression": "192*256+2*64"},
        {"policy": "AR", "payload_buffer_bytes": 256, "persistent_bytes": 52224, "expression": "192*(256+16)"},
        {"policy": "BS", "payload_buffer_bytes": 1024, "persistent_bytes": 50432, "expression": "48*(4*256+24)+2*64"},
    ])
    checks, sweep = run_tests()
    write_json(ROOT / "tests" / "test_results.json", {
        "evidence_kind": "executed local synthetic model checks",
        "named_checks": len(checks), "named_checks_passed": sum(c["pass"] for c in checks),
        "single_cut_cases": len(sweep), "checks": checks,
        "limitation": "Finite checks; not device experiments or exhaustive crash-consistency proof",
    })
    write_csv(ROOT / "tests" / "single_cut_sweep.csv", sweep)
    draw_raster()
    draw_storage()
    build_docx(result_map)
    source = (ROOT / "manuscript_source.md").read_text(encoding="utf-8")
    abstract = source.split("## Abstract\n", 1)[1].split("## 1 Introduction", 1)[0]
    body = source.split("## 1 Introduction\n", 1)[1].split("## Internal sources", 1)[0]
    body_prose = "\n\n".join(p for p in body.split("\n\n")
                              if p.strip() and not p.startswith(("##", "[[", "Figure ", "Table ")))
    with ZipFile(ROOT / "manuscript_input.docx") as z:
        xml = etree.fromstring(z.read("word/document.xml"))
        ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
              "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
        stats = {"synthetic_demo": True, "abstract_words": count_words(abstract),
                 "body_continuous_prose_words": count_words(body_prose),
                 "complete_source_words_including_headings_captions_sources": count_words(source),
                 "native_omml_equations": len(xml.xpath("//m:oMath", namespaces=ns)),
                 "native_tables": len(xml.xpath("//w:tbl", namespaces=ns)),
                 "embedded_figures": len(xml.xpath("//wp:inline", namespaces=ns)),
                 "figure_width_emu": [int(x.get("cx")) for x in xml.xpath("//wp:extent", namespaces=ns)],
                 "figure_height_emu": [int(x.get("cy")) for x in xml.xpath("//wp:extent", namespaces=ns)],
                 "result_rows": len(results), "event_rows": len(all_events),
                 "named_model_checks": len(checks), "single_cut_sweep_cases": len(sweep),
                 "word_count_rule": "ASCII alphanumeric tokens, internal apostrophes and hyphens retained; continuous prose excludes headings, captions, table, equation, abstract and sources",
                 "render_status": "not rendered by material author; parent task owns Word render inspection"}
    assert 1800 <= count_words(source) <= 2400
    assert 130 <= stats["abstract_words"] <= 170
    assert 1600 <= stats["body_continuous_prose_words"] <= 2400
    assert stats["native_omml_equations"] == 1 and stats["native_tables"] == 1 and stats["embedded_figures"] == 2
    assert stats["figure_width_emu"] == [5760000, 5760000]
    assert stats["figure_height_emu"] == [4032000, 4032000]
    assert len(results) == 36 and all(r["durable_output_records"] == 192 for r in results)
    write_json(ROOT / "input_integrity.json", stats)
    paths = sorted(p for p in ROOT.rglob("*") if p.is_file()
                   and p.name != "SHA256SUMS.txt" and "__pycache__" not in p.parts)
    (ROOT / "SHA256SUMS.txt").write_text("".join(
        f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(ROOT).as_posix()}\n" for p in paths), encoding="utf-8")
    print(json.dumps(stats, indent=2))

if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out',type=Path,required=True)
    parser.add_argument('--font',type=Path,required=True)
    parser.add_argument('--bold-font',type=Path,required=True)
    args=parser.parse_args()
    if args.out.exists(): raise ValueError('Use a new output directory; preserve the source material.')
    if not args.font.is_file() or not args.bold_font.is_file(): raise FileNotFoundError('Both font files are required.')
    args.out.mkdir(parents=True)
    for name in ['manuscript_source.md','specification.md','implementation_log.md','request.md','model.py','README.md']:
        shutil.copy2(ROOT/name,args.out/name)
    ROOT=args.out.resolve();FONT=args.font;BOLD_FONT=args.bold_font
    main()
