"""Original editing DEMO; no experimental or publication claims."""
from pathlib import Path
import argparse,json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def main():
    p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);a=p.parse_args();a.out.mkdir(parents=True,exist_ok=False)
    d=Document();first=d.add_paragraph('The update uses ')
    math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='x';r.append(t);math.append(r);first._p.append(math)
    first.add_run('. The next decision sets the buffer bound. ');first.add_run('y').italic=True
    second=d.add_paragraph('Historical text: ');ins=OxmlElement('w:ins');ins.set(qn('w:id'),'7');ins.set(qn('w:author'),'Demo author');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='retained insertion';r.append(t);ins.append(r);second._p.append(ins)
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'demo_source');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=' source';r.append(t);link.append(r);second._p.append(link)
    table=d.add_table(rows=2,cols=2);table.cell(0,0).text='DEMO';table.cell(0,1).text='Value';table.cell(1,0).text='Unchanged';table.cell(1,1).text='4'
    d.save(a.out/'input.docx')
    edits=[{'paragraph_id':'P0001','operation':'replace_span','before':'The next decision sets the buffer bound.','after':'The buffer bound is the next decision.','split_after':['x.'],'reason':'Separate the execution statement from the next design decision.','source':['Original DEMO text; no scientific result.']}]
    (a.out/'edits.json').write_text(json.dumps(edits,indent=2),encoding='utf-8')
if __name__=='__main__':main()
