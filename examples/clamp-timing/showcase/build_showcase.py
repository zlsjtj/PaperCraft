"""从公开改稿案例制作可读对照；不改论文、不生成研究数据。"""
from pathlib import Path
import argparse, hashlib, json, subprocess
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from PIL import Image
from docx import Document
from xml.sax.saxutils import escape

W,H=1280,800
INK='#203139';MUTED='#596970';ACCENT='#A44732';BLUE='#276C81';PAPER='#FAF8F3'


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--case',type=Path,required=True);p.add_argument('--clean-pages',type=Path,required=True)
    p.add_argument('--review-pages',type=Path,required=True);p.add_argument('--out',type=Path,required=True)
    p.add_argument('--font',type=Path,required=True);p.add_argument('--bold-font',type=Path,required=True)
    p.add_argument('--pdftoppm',type=Path,required=True);a=p.parse_args()
    if a.out.exists():p.error('输出目录已存在')
    a.out.mkdir(parents=True)
    pdfmetrics.registerFont(TTFont('zh',str(a.font),subfontIndex=0));pdfmetrics.registerFont(TTFont('zhb',str(a.bold_font),subfontIndex=0))
    clean=Document(a.case/'clean.docx');rough=Document(a.case/'input/rough.docx')
    assert 'When to lock' in clean.paragraphs[1].text
    assert '125 micrometres' in clean.paragraphs[3].text
    assert 'Holding contact' in clean.paragraphs[8].text
    specs=[
        ('contribution','创新定位','同一套夹具，改变的是锁定时机'),
        ('engineering','技术工作','把操作写成有理由的设计决定'),
        ('delivery','图文交付','同一条论证，进入完整的清稿与审阅稿'),
    ]
    records=[]
    for number,(name,topic,title) in enumerate(specs,1):
        c=canvas.Canvas(str(a.out/(name+'.pdf')),pagesize=(W,H),invariant=1)
        c.setTitle('PaperCraft — '+topic);svg=[];positions=[]
        def rect(x,y,w,h,col):
            c.setFillColor(HexColor(col));c.rect(x,H-y-h,w,h,fill=1,stroke=0)
            svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{col}"/>')
        def line(x,y,x2,y2,col='#D9DDD9',width=1):
            c.setStrokeColor(HexColor(col));c.setLineWidth(width);c.line(x,H-y,x2,H-y2)
            svg.append(f'<path d="M{x},{y} L{x2},{y2}" stroke="{col}" stroke-width="{width}"/>')
        def txt(t,x,y,size=26,bold=False,col=INK):
            font='zhb' if bold else 'zh';tw=pdfmetrics.stringWidth(t,font,size)
            assert x+tw<=W-30,(name,t,x+tw)
            c.setFont(font,size);c.setFillColor(HexColor(col));c.drawString(x,H-y,t)
            svg.append(f'<text x="{x}" y="{y}" font-family="Microsoft YaHei" font-weight="{700 if bold else 400}" font-size="{size}" fill="{col}">{escape(t)}</text>')
            positions.append({'text':t,'x':x,'y':y,'size':size})
        def para(t,x,y,width,size=26,leading=42,col=INK):
            current=''
            for ch in t:
                if ch in '。，；：！？、' and current and pdfmetrics.stringWidth(current+ch,'zh',size)>width:
                    txt(current+ch,x,y,size,col=col);current='';y+=leading
                elif ch=='\n' or pdfmetrics.stringWidth(current+ch,'zh',size)>width:
                    txt(current,x,y,size,col=col);current='' if ch=='\n' else ch;y+=leading
                else:current+=ch
            if current:txt(current,x,y,size,col=col);y+=leading
            return y
        def img(path,x,y,w,h):
            import base64
            c.drawImage(ImageReader(str(path)),x,H-y-h,width=w,height=h)
            encoded=base64.b64encode(path.read_bytes()).decode()
            svg.append(f'<image x="{x}" y="{y}" width="{w}" height="{h}" href="data:image/png;base64,{encoded}"/>')
        rect(0,0,W,H,PAPER)
        txt('PaperCraft',54,57,23,True);txt(f'0{number} / {topic}',990,57,21,col=MUTED)
        line(54,78,1226,78);txt(title,54,144,42,True)
        if name in ('contribution','engineering'):
            txt('英文原稿与清稿的中文摘述 · 公开教学构造案例',55,189,20,col=MUTED)
            txt('原稿',55,260,24,True,col=MUTED);txt('改稿',661,260,24,True,col=ACCENT)
            line(625,233,625,463)
            if name=='contribution':
                para('比较了已有压头的三种模式，记录包含九行结果。L 模式在斜面条件下的力变异系数为 2.3%，漂移为 19 μm。其他条件也有差异。',55,308,518,27,43,col=MUTED)
                para('压头转动有助于贴合斜面，但加载时保持自由也可能漂移。比较的关键是何时锁住：先贴合，再锁定，并为较低漂移付出额外准备时间。',661,308,557,27,43)
                txt('同一斜面，F → L：收益与代价一起出现',55,527,27,True)
                for x,label,value in [(55,'加载期间漂移','125 → 19 μm'),(467,'准备时间','17 → 32 s'),(877,'末端力 CV','2.1 → 2.3%')]:
                    txt(label,x,585,22,col=MUTED);txt(value,x,641,38,True,col=BLUE)
                txt('CV 差异不能据此判定等效或显著；平面与横向斜面结果仍在完整稿中。',55,710,21,col=MUTED)
            else:
                para('自由接近，保持预载后锁定；锁前需要预载确认，随后置零并加载。每种模式使用相同夹具和设置。',55,311,518,27,44,col=MUTED)
                para('先确认预载保持，再在不抬起压头的情况下锁定，让锁止作用于已就位姿态。锁后置零，统一加载增量的起点。',661,311,555,27,44)
                actions=[('确认预载','缺确认则终止'),('不抬起地锁定','保留已就位姿态'),('锁后置零','统一增量起点'),('加载 0.20 mm','卸载后才解锁')]
                for j,(top,bottom) in enumerate(actions):
                    x=55+j*298;line(x,523,x+255,523,BLUE,2)
                    txt(top,x,570,25,True);txt(bottom,x,614,21,col=MUTED)
                txt('必要性来自实现说明；现有记录全部收到确认，不能写成已经测过故障恢复。',55,703,21,col=MUTED)
        else:
            txt('下面是实际交付文件，网页缩图用于看结构；全文与修改请打开 PDF。',55,190,22,col=MUTED)
            pages=sorted(a.clean_pages.glob('page-*.png'))
            rp=sorted(a.review_pages.glob('page-*.png'))
            assert len(pages)==len(rp)==3
            chosen=[(pages[0],'清稿 · 问题与方法'),(pages[1],'清稿 · 机制与记录'),(rp[1],'审阅稿 · 黄色改写')]
            for j,(path,label) in enumerate(chosen):
                im=Image.open(path);w=332;h=w*im.height/im.width;x=55+j*409;y=235
                rect(x-1,y-1,w+2,h+2,'#D6DDD9');img(path,x,y,w,h);txt(label,x,720,22,True)
            txt('公式、九行表值、引用与原有标记保留；黄色表示文字改动，不是原生修订。',55,762,20,col=MUTED)
        if name!='delivery':
            line(54,743,1226,743);txt('来源：input/rough.docx · input/implementation.md · clean.docx · input/results.csv',55,776,17,col=MUTED)
        c.showPage();c.save()
        (a.out/(name+'.svg')).write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="800" viewBox="0 0 1280 800"><title>{escape(title)}</title>'+''.join(svg)+'</svg>',encoding='utf-8')
        subprocess.run([str(a.pdftoppm),'-png','-singlefile','-r','72',str(a.out/(name+'.pdf')),str(a.out/name)],check=True,capture_output=True)
        records.append({'name':name,'title':title,'text':positions})
    # 幻灯导览来自实际案例；不是模型实时生成的屏幕录制。
    slides=[Image.open(a.out/(s[0]+'.png')).convert('RGB').resize((960,600),Image.Resampling.LANCZOS) for s in specs]
    slides[0].save(a.out/'tour.gif',save_all=True,append_images=slides[1:],duration=[11000,11000,8000],loop=0,optimize=True)
    proof={'case_sources':{rel:hashlib.sha256((a.case/rel).read_bytes()).hexdigest() for rel in ['input/rough.docx','input/implementation.md','input/results.csv','clean.docx','review.docx']},'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'content':'Chinese abridged presentation of existing English revisions; no new experiment or paper rewrite','generation':'fixed showcase export, not a fresh skill generation test','slides':records,'tour':'30-second edited case tour; not a live generation recording'}
    (a.out/'record.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':str(a.out),'slides':3}))


if __name__=='__main__':main()
