"""核对真实交付包里的受保护对象、数值和图片；不评价论文说服力。"""
from pathlib import Path
import argparse, csv, hashlib, json, zipfile
from lxml import etree as E
from docx import Document
from pypdf import PdfReader

N={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
sha=lambda b:hashlib.sha256(b).hexdigest()
def root(p):
    with zipfile.ZipFile(p) as z:return E.fromstring(z.read('word/document.xml'))
def nodes(r,x):return [E.tostring(n,method='c14n',exclusive=True) for n in r.xpath(x,namespaces=N)]
def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--delivery',type=Path,required=True);p.add_argument('--figures',type=Path,required=True);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    case=Path(__file__).resolve().parent.parent;source=case/'input/rough.docx';s=root(source);rows=list(csv.DictReader((case/'input/results.csv').open(encoding='utf-8-sig')))
    expected=[[r[k] for k in ['condition','mode','n_reseatings','F_cv_percent','Drift_um','Seat_s']] for r in rows]
    record={'source_sha256':sha(source.read_bytes()),'checks':{},'pdf_checks':{},'visual_review':'separately recorded model inspection','human_approval':False}
    for label in ['clean','review']:
        path=a.delivery/(label+'.docx');d=Document(path);r=root(path)
        checks={}
        for name,x in [('formula','.//m:oMath'),('fields','.//w:fldChar|.//w:instrText'),('bookmarks','.//w:bookmarkStart|.//w:bookmarkEnd'),('hyperlinks','.//w:hyperlink'),('historical_revisions','.//w:ins|.//w:del|.//w:moveFrom|.//w:moveTo')]:checks[name]=nodes(s,x)==nodes(r,x)
        checks['italic_variables']=nodes(s,'.//w:r[w:rPr/w:i]')==nodes(r,'.//w:r[w:rPr/w:i]')
        checks['original_yellow_retained']=all(v in nodes(r,".//w:r[w:rPr/w:highlight[@w:val='yellow']]") for v in nodes(s,".//w:r[w:rPr/w:highlight[@w:val='yellow']]"))
        checks['table_cell_values']=[[c.text for c in row.cells] for row in d.tables[0].rows[1:]]==expected
        checks['table_header']=[c.text for c in d.tables[0].rows[0].cells]==[c.text for c in Document(source).tables[0].rows[0].cells]
        checks['two_figures_at_160mm']=len(d.inline_shapes)==2 and all(abs(s.width/36000-160)<.01 for s in d.inline_shapes)
        with zipfile.ZipFile(path) as z:
            media={sha(z.read(n)) for n in z.namelist() if n.startswith('word/media/')}
        checks['exact_png_svg_embedded']=all(sha((a.figures/n).read_bytes()) in media for n in ['mechanism.png','mechanism.svg','results.png','results.svg'])
        pdf=a.delivery/(label+'.pdf')
        record['pdf_checks'][label]='RUN' if pdf.exists() else 'NOT_RUN'
        if pdf.exists():
            pages=PdfReader(pdf).pages;checks['three_pdf_pages']=len(pages)==3
            checks['page_roles']=all(word in pages[i].extract_text() for i,word in enumerate(['Figure 1.','Table 1.','Figure 2.']))
        record['checks'][label]=checks
    record['checks']['same_prose']= [p.text for p in Document(a.delivery/'clean.docx').paragraphs]==[p.text for p in Document(a.delivery/'review.docx').paragraphs]
    # 从实际 SVG 图元重算横坐标，而不只核对生成器自报的坐标表。
    svg=E.parse(str(a.figures/'results.svg'));S={'s':'http://www.w3.org/2000/svg'};actual=[]
    for m,tag in [('E','polygon'),('F','circle'),('L','polygon')]:
        for el in svg.findall('.//s:'+tag,S):
            if tag=='circle':x=float(el.get('cx'));y=float(el.get('cy'))
            else:
                pts=[[float(v) for v in p.split(',')] for p in el.get('points').split()]
                if m=='E':
                    if len(pts)!=4 or abs(pts[1][0]-pts[0][0]-8)>1e-9:continue
                    x=pts[0][0]+4;y=pts[0][1]+4
                else:
                    if len(pts)!=3:continue
                    x=pts[0][0];y=pts[0][1]+6
            if y<100:continue
            ci=min(range(3),key=lambda i:abs(y-(132+107*i+(['E','F','L'].index(m)-1)*22)))
            col=min(range(3),key=lambda i:abs(x-([205,484,763][i]+103)))
            key,maxv=[('F_cv_percent',8),('Drift_um',140),('Seat_s',35)][col]
            value=(x-[205,484,763][col])/206*maxv
            condition=['flat','pitch_1deg','roll_1deg'][ci]
            row=next(r for r in rows if r['mode']==m and r['condition']==condition)
            assert abs(value-float(row[key]))<1e-7,(m,condition,key,value)
            actual.append((m,condition,key))
    record['checks']['actual_svg_27_values']=len(actual)==len(set(actual))==27
    record['status']='PASS' if all(all(v.values()) if isinstance(v,dict) else v for v in record['checks'].values()) else 'FAIL'
    if record['status']=='PASS' and 'NOT_RUN' in record['pdf_checks'].values():record['status']='WORD_AND_SVG_PASS_PDF_NOT_RUN'
    record['note']='表格内容与行列保留；排版属性有变化，未声称表 XML 完全相同。新图为两个 SVG/PNG 对应表示。'
    a.out.write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(record,ensure_ascii=False))
    if record['status']=='FAIL':raise SystemExit(1)
if __name__=='__main__':main()
