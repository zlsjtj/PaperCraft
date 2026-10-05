"""组装公开夹具案例的完整图文稿，保留原始公式、表值、域和历史黄色。"""
from pathlib import Path
import argparse, copy, hashlib, json, shutil, subprocess, sys, zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree as E

HERE=Path(__file__).resolve().parent; CASE=HERE.parent; ROOT=CASE.parents[1]
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
CAP1=('Figure 1. Use the same pitch freedom to seat the pad, then lock it for loading. '
      'The two views show mode L: seat at 10 N, hold and acknowledge preload, close the clamp without lifting, '
      'then zero and ramp 0.20 mm. Lock symbols denote clamp state, not internal construction. '
      'The strip compares E, F and L; open and filled circles mean free and locked. '
      'The pin permits only pitch. Perspective and slope are schematic, not measured geometry or pressure maps.')
CAP2=('Figure 2. Complete constructed comparison. Squares, circles and triangles denote E, F and L. '
      'Each row group is one coupon condition; horizontal position gives the metric and vertical offsets only separate modes. '
      'All nine records and all three metrics are shown. Each record summarizes eight reseatings of one coupon, '
      'not eight independent specimens. No uncertainty intervals are available; exact values are in Table 1.')

def plan():
    d=json.loads((CASE/'authored-edits.json').read_text('utf-8'))
    new={
      'P0004':('On a sloping coupon, the freedom that lets a compression pad seat also remains available during loading. '
               'We compare when to lock an existing single-axis pad: before approach (E), remain free (F), or seat under preload and then lock (L). '
               'In the constructed 1 degree pitch-slope records, L trades 15 s of extra preparation for lower median ramp drift '
               '(19 versus 125 micrometres for F). The endpoint force coefficient of variation (CV) is 2.3% versus 2.1%, '
               'without uncertainty estimates; early locking gives 7.2%. '
               'Flat coupons gain no repeatability, and a cross-axis slope remains uncompensated. '
               'The comparison identifies a conditional benefit of changing the operating sequence on the same hardware.'),
      'P0007':('We vary the timing of one existing friction clamp while retaining the frame, sensors, calibration, pad surface and analysis. '
               'Figure 1 separates the freedom needed for seating from the restraint applied during loading. '
               'A pitch slope tests this sequence where the pivot can accommodate the surface; flat and roll-sloped coupons test its limits. '
               'This operating-sequence comparison uses constructed records [1,2], with no new pivot, controller or feedback algorithm.'),
      'P0009':('Mode E fixes the horizontal pad before approach to 10 N. Mode F approaches with the clamp open and leaves it open during the ramp. '
               'Mode L approaches freely to 10 N, holds that preload, and locks without lifting the pad. '
               'Holding contact makes the lock act on the seated pose rather than requiring a second approach. '
               'A preload-hold acknowledgement is required before locking; its absence aborts the run. All supplied records contain it. '
               'The displacement increment is zeroed after the prescribed approach or lock sequence, so every mode then receives the same 0.20 mm ramp. '
               'The clamp is released only after unloading.'),
      'P0017':('The pitch rows show why the comparison needs both E and F. E limits drift to 12 micrometres but gives 7.2% force CV. '
               'F gives 2.1% CV but 125 micrometres drift; L gives 2.3% CV and 19 micrometres drift. '
               'Relative to F, L therefore trades 15 additional seconds of preparation for 106 micrometres less median drift. '
               'The small CV difference does not establish equivalence or significance without uncertainty estimates. '
               'Figure 2 keeps this tradeoff alongside the other conditions.\n\n'
               'On flat tops, force CV spans only 1.3-1.4%. L takes 31 s rather than 18 s for E, with drift of 16 rather than 14 micrometres; '
               'the extra action offers no repeatability gain on these records. On the roll slope, all modes remain at 6.8-7.0% CV. '
               'L reduces drift relative to F but does not repair this variation. The useful sequence is therefore conditional on the available pitch freedom, '
               'not a general remedy for a sloping surface.'),
      'P0019':CAP1,
      'P0021':('The comparison turns an existing degree of freedom into a choice about timing: use it to seat, then restrain it for loading when '
               'the reduction in drift justifies the additional preparation. Flat and roll-sloped coupons prevent a universal preference for L. '
               'Repeated mountings establish neither force accuracy, material properties nor pressure uniformity. '
               'These constructed records demonstrate how to assess an operating sequence; they do not validate physical fixture performance.')}
    for e in d['edits']:
        if e['paragraph_id'] in new:
            e['after']=new[e['paragraph_id']]
            e['reason']='Connect the same operating choice across entry, necessary steps, mechanism and complete evidence; retain its cost and directional limits.'
    return d

def add_svg(doc,png_run,svg_path,docx_path):
    # 在现有 PNG drawing 上绑定可编辑 SVG 备选；两种表示来自同一生成稿。
    with zipfile.ZipFile(docx_path) as z:parts={n:z.read(n) for n in z.namelist()}
    root=E.fromstring(parts['word/document.xml']);rr=E.fromstring(parts['word/_rels/document.xml.rels']);ct=E.fromstring(parts['[Content_Types].xml'])
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    for i,blip in enumerate(root.findall('.//a:blip',ns),1):
        extlist=E.SubElement(blip,'{'+ns['a']+'}extLst');ext=E.SubElement(extlist,'{'+ns['a']+'}ext',uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        rid='rIdStorySVG'+str(i);part='story'+str(i)+'.svg'
        E.SubElement(ext,'{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',{'{'+ns['r']+'}embed':rid})
        E.SubElement(rr,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=rid,Type=ns['r']+'/image',Target='media/'+part)
        parts['word/media/'+part]=(svg_path/('mechanism.svg' if i==1 else 'results.svg')).read_bytes()
    if not any(n.get('Extension')=='svg' for n in ct):E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='svg',ContentType='image/svg+xml')
    for k,v in [('word/document.xml',root),('word/_rels/document.xml.rels',rr),('[Content_Types].xml',ct)]:parts[k]=E.tostring(v,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(docx_path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)

def assemble(src,out,figures,label):
    doc=Document(src);normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(7)
    for sec in doc.sections:
        sec.page_width=Mm(210);sec.page_height=Mm(297);sec.left_margin=Mm(25);sec.right_margin=Mm(25);sec.top_margin=Mm(18);sec.bottom_margin=Mm(18)
    for name in ['Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(21 if name=='Title' else 12)
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(7)
        pp=st.element.find(qn('w:pPr'))
        if pp is not None:
            for n in list(pp):
                if n.tag==qn('w:pBdr'):pp.remove(n)
    for p in doc.paragraphs:
        p.paragraph_format.keep_with_next=False
        for border in list(p._p.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
        if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
    title=next(p for p in doc.paragraphs if p.text.startswith('When to lock'));title.style='Title'
    title.paragraph_format.keep_with_next=True
    method=next(p for p in doc.paragraphs if p.text=='Method');method.paragraph_format.page_break_before=True
    result=next(p for p in doc.paragraphs if p.text=='Results');result.paragraph_format.page_break_before=True
    # 原稿这个段落只含占位符，验证后才替换；未清空带公式或域的段落。
    pic=next(p for p in doc.paragraphs if p.text=='[Figure 1 placeholder]')
    assert not pic._p.xpath('.//m:oMath|.//w:fldChar|.//w:hyperlink|.//w:ins|.//w:del')
    pic.clear();pic.add_run().add_picture(str(figures/'mechanism.png'),width=Mm(160))
    cap=next(p for p in doc.paragraphs if p.text==CAP1)
    method._p.addprevious(pic._p);method._p.addprevious(cap._p)
    pic.paragraph_format.keep_with_next=True
    # 在结果导读之后插入全量定量图，图注与主文各承担自己的解释。
    lead=next(p for p in doc.paragraphs if p.text.startswith('All records'))
    plot=doc.add_paragraph();plot.add_run().add_picture(str(figures/'results.png'),width=Mm(160));lead._p.addnext(plot._p)
    cap2=doc.add_paragraph(CAP2);plot._p.addnext(cap2._p);plot.paragraph_format.keep_with_next=True
    for p in [cap,cap2]:
        p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(9)
    if label=='review':
        for r in cap2.runs:r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    # 表格保持原值及行列，移到方法页，读者仍可逐项核查。
    tbl=doc.tables[0];tc=next(p for p in doc.paragraphs if p.text.startswith('Table 1.'))
    result._p.addprevious(tc._p);result._p.addprevious(tbl._tbl);tc.paragraph_format.keep_with_next=True
    widths=[31,16,25,31,28,29];tbl.autofit=False
    for rowi,row in enumerate(tbl.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Mm(widths[ci]);props=cell._tc.get_or_add_tcPr()
            for t in ['tcBorders','shd','tcMar']:
                for n in list(props):
                    if n.tag==qn('w:'+t):props.remove(n)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                edge_node=OxmlElement('w:'+edge);edge_node.set(qn('w:val'),'single');edge_node.set(qn('w:sz'),'4');edge_node.set(qn('w:color'),'D9D9D9');borders.append(edge_node)
            props.append(borders)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E5EAED' if rowi==0 else ('F4F6F7' if (rowi-1)//3%2 else 'FFFFFF'));props.append(shd)
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
            props.append(mar)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci<2 else WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:r.font.size=Pt(9);r.font.name='Arial';r.bold=rowi==0
    # 固定原表格网格宽度，避免继承等宽布局。
    for n,w in zip(tbl._tbl.tblGrid.gridCol_lst,widths):n.set(qn('w:w'),str(round(w/25.4*1440)))
    path=out/(label+'.docx');doc.save(path);add_svg(doc,None,figures,path)

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--out',type=Path,required=True);p.add_argument('--figures',type=Path,required=True);a=p.parse_args()
    if a.out.exists():p.error('输出目录必须不存在')
    a.out.mkdir(parents=True);edit=a.out/'authored-edits.json';edit.write_text(json.dumps(plan(),ensure_ascii=False,indent=2),encoding='utf-8')
    subprocess.run([sys.executable,str(ROOT/'scripts/apply_authored_edits.py'),'--source',str(CASE/'input/rough.docx'),'--edits',str(edit),'--skill',str(ROOT),'--out',str(a.out/'prose')],check=True)
    for label in ['clean','review']:assemble(a.out/'prose'/(label+'.docx'),a.out,a.figures,label)
    (a.out/'captions.md').write_text(CAP1+'\n\n'+CAP2+'\n',encoding='utf-8')
    (a.out/'build.json').write_text(json.dumps({'source_sha256':sha(CASE/'input/rough.docx'),'builder_sha256':sha(__file__),'papercraft_skill_sha256':sha(ROOT/'SKILL.md'),'figure_sources':{n:sha(a.figures/n) for n in ['mechanism.svg','mechanism.png','results.svg','results.png']},'changes':'Narrative and figure sequence; exact original table cell texts retained with new formatting.','yellow':'Relative to rough.docx; original highlighted statement also retained. Not native Track Changes.','visual_review':'PENDING'},indent=2),encoding='utf-8')
    print(a.out)

if __name__=='__main__':main()
