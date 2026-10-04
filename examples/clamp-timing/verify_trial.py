from pathlib import Path
import argparse, csv, hashlib, json, re, zipfile, subprocess, sys
from lxml import etree as E
from docx import Document
from pypdf import PdfReader

N={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def xml(p):
    with zipfile.ZipFile(p) as z:return E.fromstring(z.read('word/document.xml'))
def sig(n):return E.tostring(n,method='c14n',exclusive=True)
def nodes(root,x):return [sig(n) for n in root.xpath(x,namespaces=N)]
def main():
    ap=argparse.ArgumentParser();ap.add_argument('stage',type=Path);ap.add_argument('--papercraft',type=Path,required=True);a=ap.parse_args();s=a.stage
    source=s/'inputs/rough.docx';src=xml(source);rows=list(csv.DictReader((s/'inputs/results.csv').open(encoding='utf-8-sig')))
    result={'source_sha256':sha(source),'checks':{},'limitations':['No Word rendering in child; parent has exclusive renderer role.','Automated figure relation endpoints and logical count scope remain REVIEW_REQUIRED; human author acceptance absent.']}
    old_yellow=nodes(src,".//w:r[w:rPr/w:highlight[@w:val='yellow']]")
    old_italic=nodes(src,".//w:r[w:rPr/w:i]")
    for label in ['clean','review']:
        path=s/(label+'.docx');r=xml(path);doc=Document(path)
        actual=[[c.text for c in row.cells] for row in doc.tables[0].rows[1:]]
        expected=[[d['condition'],d['mode'],d['n_reseatings'],d['F_cv_percent'],d['Drift_um'],d['Seat_s']] for d in rows]
        checks={
          'original_omml_xml':nodes(src,'.//m:oMath')==nodes(r,'.//m:oMath'),
          'table_xml_unchanged':nodes(src,'.//w:tbl')==nodes(r,'.//w:tbl'),
          'all_nine_csv_rows_exact':actual==expected and len(actual)==9,
          'italic_runs_exact':old_italic==nodes(r,".//w:r[w:rPr/w:i]"),
          'ref_instruction_exact':nodes(src,'.//w:instrText')==nodes(r,'.//w:instrText'),
          'field_nodes_exact':nodes(src,'.//w:fldChar')==nodes(r,'.//w:fldChar'),
          'bookmark_nodes_exact':nodes(src,'.//w:bookmarkStart|.//w:bookmarkEnd')==nodes(r,'.//w:bookmarkStart|.//w:bookmarkEnd'),
          'hyperlink_xml_exact':nodes(src,'.//w:hyperlink')==nodes(r,'.//w:hyperlink'),
          'existing_yellow_run_retained':all(x in nodes(r,".//w:r[w:rPr/w:highlight[@w:val='yellow']]") for x in old_yellow),
          'exactly_one_new_drawing':len(r.xpath('.//w:drawing',namespaces=N))==1,
        }
        result['checks'][label]=checks
        body=' '.join(p.text for p in doc.paragraphs if not p.text.startswith(('CONSTRUCTED','Figure 1.','Table 1.','Sources','[1]')))
        result[label]={'sha256':sha(path),'body_word_count_regex':len(re.findall(r'\b\S+\b',body)),'highlight_run_count':len(nodes(r,'.//w:r[w:rPr/w:highlight]'))}
        sys.path.insert(0,str(a.papercraft/'scripts'));import audit_figure_integration as afi
        d=afi.inspect(path)['drawings'][0]
        reps=[{'kind':rep['kind'],'media_sha256':rep['sha256'],'source_asset':'figure/figure.svg' if rep['kind']=='svg' else 'figure/figure.png','source_asset_sha256':rep['sha256']} for rep in d['representations']]
        manifest={'scope':'all-main-document-drawings','docx_sha256':sha(path),'figures':[{'id':'Figure 1','drawing_index':1,'disposition':'revised','media_sha256':d['sha256'],'placement_width_mm':160,'placement_height_mm':100,'caption_contains':'Figure 1.','caption_exact':(s/'caption.txt').read_text(encoding='utf-8'),'source_asset':'figure/figure.png','source_asset_sha256':sha(s/'figure/figure.png'),'representations':reps}]}
        manifest_path=s/(label+'-integration-manifest.json');manifest_path.write_text(json.dumps(manifest,indent=2),encoding='utf-8')
        audit=afi.audit(path,manifest_path,s/(label+'-bound.docx'))
        (s/(label+'-integration-audit.json')).write_text(json.dumps(audit,indent=2),encoding='utf-8')
        checks['png_and_svg_hash_bound_at_160x100']=audit['technical_status']=='PASS'
    result['checks']['clean_review_visible_text_equal']=[p.text for p in Document(s/'clean.docx').paragraphs]==[p.text for p in Document(s/'review.docx').paragraphs]
    page=PdfReader(s/'figure/figure.pdf').pages[0]
    result['figure_pdf_mm']=[round(float(page.mediabox.width)*25.4/72,5),round(float(page.mediabox.height)*25.4/72,5)]
    result['figure_text_editable']=len(E.parse(str(s/'figure/figure.svg')).findall('.//{http://www.w3.org/2000/svg}text'))>0
    result['status']='PASS' if all(all(v.values()) if isinstance(v,dict) else v for v in result['checks'].values()) else 'FAIL'
    (s/'fidelity-checks.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps({'stage':str(s),'status':result['status'],'words':result['clean']['body_word_count_regex'],'figure_mm':result['figure_pdf_mm']}))
if __name__=='__main__':main()
