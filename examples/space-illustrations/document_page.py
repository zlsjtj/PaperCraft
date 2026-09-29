"""Build inspectable one-page DEMO documents, without changing any manuscript."""
from pathlib import Path
import argparse, hashlib, json, zipfile
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from lxml import etree

def build(out, title, text, caption, image, review=False, highlight_caption=True):
    if out.exists():
        raise FileExistsError(out)
    d=Document();s=d.sections[0]
    s.page_width=Mm(210);s.page_height=Mm(297)
    s.top_margin=Mm(23);s.bottom_margin=Mm(23)
    s.left_margin=Mm(25);s.right_margin=Mm(25)
    for name in ['Normal','Title','Caption']:
        st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor.from_string('25333B');st.font.bold=False;st.font.italic=False
        pp=st.element.find(qn('w:pPr'))
        if pp is not None:
            for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
    d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.line_spacing=1.12
    d.styles['Normal'].paragraph_format.space_after=Pt(9)
    d.styles['Title'].font.size=Pt(18)
    d.styles['Title'].paragraph_format.space_after=Pt(12)
    d.styles['Caption'].font.size=Pt(9.5)
    d.styles['Caption'].paragraph_format.line_spacing=1.1
    d.add_paragraph(title,'Title')
    p=d.add_paragraph();r=p.add_run(text)
    if review:r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    p=d.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(image),width=Mm(160))
    p=d.add_paragraph(style='Caption');r=p.add_run('Figure 1. '+caption)
    if review and highlight_caption:r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    p=d.add_paragraph('DEMO — fictional structure; no experimental results.',style='Caption')
    if review:d.add_paragraph('Yellow marks changed paragraphs or captions at paragraph level. The figure was replaced; this is not native Word Track Changes.',style='Caption')
    d.core_properties.title=title;d.core_properties.subject='Original scientific illustration development DEMO'
    d.save(out)
    with zipfile.ZipFile(out) as z:
        root=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
        media=[z.read(n) for n in z.namelist() if n.startswith('word/media/')]
        assert len(media)==1 and media[0]==image.read_bytes()
        assert int(root.find('.//wp:extent',ns).get('cx'))==5760000
        assert text in ''.join(root.xpath('//w:t/text()',namespaces=ns))
        assert not root.xpath('//w:ins|//w:del',namespaces=ns)
        count=len(root.xpath('//w:highlight',namespaces=ns));assert count==((1+int(highlight_caption)) if review else 0)
    return {'file':out.name,'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'embedded_image_sha256':hashlib.sha256(image.read_bytes()).hexdigest(),'width_mm':160,'text_verified':True,'highlighted_runs':count,'native_revisions':0,'render_status':'PENDING'}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--spec',required=True,type=Path);ap.add_argument('--out',required=True,type=Path);a=ap.parse_args()
    spec=json.loads(a.spec.read_text(encoding='utf8'));a.out.mkdir(parents=True,exist_ok=True)
    results=[]
    for c in spec['cases']:
        for mode in ['before','clean','review']:
            old=mode=='before'
            results.append(build(a.out/(c['id']+'-'+mode+'.docx'),c['title'],c['before'] if old else c['after'],c['before_caption'] if old else c['caption'],Path(c['before_image'] if old else c['image']),mode=='review',c['before_caption']!=c['caption']))
    (a.out/'document-audit.json').write_text(json.dumps(results,indent=2),encoding='utf8')
    print(json.dumps(results,indent=2))
if __name__=='__main__':main()
