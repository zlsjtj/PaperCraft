"""Generate a self-contained mixed-format Word fixture and a local span plan."""
from pathlib import Path
import argparse,hashlib,json,sys,base64,io
from docx import Document
from docx.oxml import OxmlElement as X
from docx.oxml.ns import qn
from docx.shared import Inches

def build(out):
    out=Path(out);out.mkdir(parents=True,exist_ok=False)
    d=Document();d.add_heading('Thermal model teaching example',0)
    d.add_paragraph('Constructed arithmetic example. It is not a physical experiment.')
    p=d.add_paragraph();p.add_run('The design proves better cooling. ');p.add_run('T').italic=True;p.add_run(' is defined by ')
    m=X('m:oMath');r=X('m:r');t=X('m:t');t.text='T = R × P';r.append(t);m.append(r);p._p.append(m);p.add_run('; see ')
    for kind in ['begin','instruction','separate','result','end']:
        r=X('w:r')
        if kind=='instruction':t=X('w:instrText');t.text=' REF ThermalTable \\h '
        elif kind=='result':t=X('w:t');t.text='Table 1'
        else:t=X('w:fldChar');t.set(qn('w:fldCharType'),kind)
        r.append(t);p._p.append(r)
    p.add_run(' and ');rid=d.part.relate_to('https://example.org/thermal-demo','http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=X('w:hyperlink');h.set(qn('r:id'),rid);r=X('w:r');t=X('w:t');t.text='the model note';r.append(t);h.append(r);p._p.append(h);p.add_run('. ')
    p.add_run('Measurement pending.').font.highlight_color=4
    for kind,tag,value in [('ins','t',' Prior insertion.'),('del','delText',' Obsolete statement.')]:
        n=X('w:'+kind);n.set(qn('w:id'),'17' if kind=='ins' else '18');n.set(qn('w:author'),'Earlier reviewer');n.set(qn('w:date'),'2026-01-01T00:00:00Z');r=X('w:r');t=X('w:'+tag);t.text=value;r.append(t);n.append(r);p._p.append(n)
    q=d.add_paragraph('Table 1 Given quantities');a=X('w:bookmarkStart');a.set(qn('w:id'),'1');a.set(qn('w:name'),'ThermalTable');b=X('w:bookmarkEnd');b.set(qn('w:id'),'1');q._p.insert(0,a);q._p.append(b)
    tab=d.add_table(rows=2,cols=2);tab.style='Table Grid'
    for c,s in zip(tab.rows[0].cells,['R','P']):c.text=s
    for c,s in zip(tab.rows[1].cells,['2 K/W','3 W']):c.text=s
    # A real embedded PNG without an image-library dependency.
    data=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
    d.add_picture(io.BytesIO(data),width=Inches(.12));d.add_paragraph('Unchanged image package sentinel, not a scientific figure.')
    doc=out/'input.docx';d.save(doc)
    sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'scripts'))
    from review_docx import inventory
    inv=inventory(doc);target=inv['paragraphs'][2]
    plan={'input_sha256':hashlib.sha256(doc.read_bytes()).hexdigest(),'scope':'local','operations':[{'id':'C1','kind':'replace_span','target':target['id'],'expect':target['text'],'spans':[{'before':'The design proves better cooling.','after':'The model gives a temperature rise from the stated resistance and power.'}],'purpose':'区分模型计算与实物散热测量','evidence':'给定模型和表1；没有物理测量','confirm':False}],'report':{'summary':'只修正一句过强表述，公式、变量、引用及历史修订保留。','limitations':['没有新增物理测量或证明散热改善。'],'checks':[]}}
    (out/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8');return doc,plan
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);a=p.parse_args();build(a.out)
