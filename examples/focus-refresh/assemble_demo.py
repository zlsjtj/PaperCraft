"""Neutral DOCX assembly for the controlled DEMO; makes no editorial choices.

Figure locations, captions, prose and tables come from the chosen manuscript.
Review highlighting marks rewritten paragraphs against the supplied draft, not
native Word revisions. Existing input files are never edited.
"""
from pathlib import Path
import argparse,re,json,hashlib
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def plain(s): return re.sub(r'[*`]', '',s).strip()
def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        # Optional placement markers are not part of the scholarly prose.
        if re.fullmatch(r'\[TABLE[_ ]\d+\]',s):i+=1;continue
        if s.startswith('$$') and s.endswith('$$') and len(s)>4:
            yield ('equation',s[2:-2]);i+=1;continue
        if s.startswith('```') or s in ('$$',r'\['):
            end='```' if s.startswith('```') else ('$$' if s=='$$' else r'\]');a=[];i+=1
            while i<len(lines) and not lines[i].strip().startswith(end): a.append(lines[i]);i+=1
            yield ('equation',' '.join(a));i+=1;continue
        if s.startswith('|'):
            a=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',v or '-') for v in row):a.append(row)
                i+=1
            yield ('table',a);continue
        if s.startswith('#'):
            yield ('heading',(len(s)-len(s.lstrip('#')),s.lstrip('#').strip()));i+=1;continue
        m=re.fullmatch(r'\[FIGURE[_ ]([123])\]',s)
        if m:yield ('figure',int(m.group(1)));i+=1;continue
        if s.startswith('!['):
            m=re.search(r'(?:figure|fig)[_ -]?([123])',s,re.I)
            # The transfer example has one image called figure.png.
            if m:yield ('figure',int(m.group(1)));i+=1;continue
            if re.search(r'\]\([^)]*figure\.png\)',s):yield ('figure',1);i+=1;continue
        a=[s];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||\[FIGURE|!\[|```|\$\$)',lines[i].strip()):a.append(lines[i].strip());i+=1
        yield ('paragraph',' '.join(a))

def runs(p,text,yellow=False):
    for part in re.split(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`)',text):
        if not part:continue
        bold=part.startswith('**');ital=part.startswith('*') and not bold
        value=part.strip('*`')
        # Render explicitly indexed source variables as typographic indices.
        # Upper-case policy identifiers and file names remain literal text.
        indexed=re.compile(r'\b([a-z]|error|T)_([A-Za-z0-9]{1,8})\b')
        pos=0
        def add(v,sub=False,variable=False):
            r=p.add_run(v);r.bold=bold;r.italic=ital or variable
            if sub:r.font.subscript=True
            if yellow:r.font.highlight_color=WD_COLOR_INDEX.YELLOW
        for m in indexed.finditer(value):
            if m.start()>pos:add(value[pos:m.start()])
            add(m[1],variable=True);add(m[2],sub=True,variable=True);pos=m.end()
        if pos<len(value):add(value[pos:])

def math_run(s):
    r=OxmlElement('m:r')
    if re.fullmatch(r'[A-Za-z]{2,}',s):
        pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
    t=OxmlElement('m:t');t.text=s;r.append(t);return r
def linear_math(p,s):
    # Native OMML runs and sub/superscripts for the short linear source equations.
    # No symbolic rearrangement or simplification is performed.
    s=s.strip('$').replace(r'\mu','μ').replace(r'\leq','≤').replace(r'\le','≤').replace(r'\times','×').replace(r'\cdot','·')
    tag=re.search(r'\\tag\{([^}]+)\}',s);s=re.sub(r'\\tag\{[^}]+\}','',s)
    s=re.sub(r'\\(?:mathrm|text)\{([^}]+)\}',r'\1',s).replace(r'\qquad','    ').replace(r'\quad','  ').replace(r'\,',' ').replace('\\ ',' ')
    if '\\' in s:raise ValueError('Unconverted math syntax requires explicit assembly repair: '+s)
    obj=OxmlElement('m:oMath')
    pattern=re.compile(r'([A-Za-z]+)(?:_(\{[^}]+\}|[A-Za-z0-9]))?(?:\^(\{[^}]+\}|[A-Za-z0-9]))?')
    pos=0
    for m in pattern.finditer(s):
        if m.start()>pos:obj.append(math_run(s[pos:m.start()]))
        base,sub,sup=m.groups()
        if sub or sup:
            n=OxmlElement('m:sSubSup' if sub and sup else ('m:sSub' if sub else 'm:sSup'));e=OxmlElement('m:e');e.append(math_run(base));n.append(e)
            for name,val in [('sub',sub),('sup',sup)]:
                if val is not None:
                    q=OxmlElement('m:'+name);q.append(math_run(val.strip('{}')));n.append(q)
            obj.append(n)
        else:obj.append(math_run(base))
        pos=m.end()
    if pos<len(s):obj.append(math_run(s[pos:]))
    p._p.append(obj)
    if tag:p.add_run('  ('+tag[1]+')')

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--manuscript',type=Path,required=True);ap.add_argument('--baseline',type=Path,required=True);ap.add_argument('--figure',action='append',required=True);ap.add_argument('--out',type=Path,required=True);a=ap.parse_args()
    if a.out.exists():raise SystemExit('New output directory required')
    a.out.mkdir(parents=True);figs={int(v.split('=',1)[0]):Path(v.split('=',1)[1]) for v in a.figure}
    text=a.manuscript.read_text(encoding='utf-8');baseline=a.baseline.read_text(encoding='utf-8');old={plain(v) for k,v in blocks(baseline) if k=='paragraph'}
    receipt={'manuscript_sha256':sha(a.manuscript),'baseline_sha256':sha(a.baseline),'figure_sha256':{str(k):sha(v) for k,v in figs.items()},'review':'yellow rewritten paragraphs; not native revisions','edits':[]}
    for mode in ['clean','review']:
        d=Document();sec=d.sections[0];sec.page_width=Mm(215.9);sec.page_height=Mm(279.4);sec.left_margin=sec.right_margin=Mm(27.95);sec.top_margin=Mm(18);sec.bottom_margin=Mm(19)
        st=d.styles['Normal'];st.font.name='Cambria';st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
        for name,size in [('Title',17),('Heading 1',13),('Heading 2',11.5),('Heading 3',10.5)]:
            st=d.styles[name];st.font.name='Cambria';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
        for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
            st=d.styles[name]
            for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
            rf=st.element.get_or_add_rPr().rFonts
            if rf is not None:
                for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                    rf.attrib.pop(qn('w:'+attr),None)
                for attr in ['ascii','hAnsi','eastAsia']:rf.set(qn('w:'+attr),'Cambria')
        f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('DEMO · ').font.size=Pt(8);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
        for k,v in blocks(text):
            if k=='heading':
                level,t=v;p=d.add_paragraph(style='Title' if level==1 else f'Heading {min(level-1,3)}');runs(p,t,mode=='review' and t not in baseline)
            elif k=='paragraph':
                p=d.add_paragraph();changed=plain(v) not in old;runs(p,v,mode=='review' and changed)
                if re.match(r'^(?:Figure\s+[123]|Fig\.\s*[123]|Table\s*[12])\.\s',plain(v),re.I):
                    for r in p.runs:r.font.size=Pt(9.5)
                    p.paragraph_format.keep_together=True
                    if plain(v).startswith('Table'):p.paragraph_format.keep_with_next=True
                if mode=='clean' and changed:receipt['edits'].append({'text':plain(v),'kind':'rewritten paragraph'})
            elif k=='equation':
                p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;linear_math(p,v);p.paragraph_format.keep_together=True
            elif k=='figure':
                p=d.add_paragraph();p.add_run().add_picture(str(figs[v]),width=Mm(160));p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
            elif k=='table':
                width=max(map(len,v));t=d.add_table(rows=len(v),cols=width);t.autofit=False
                colwidth=[160/width]*width
                if width==7 and v[0][0]=='Policy' and 'Cycle min' in v[0][-1]:
                    colwidth=[17,18,18,25,23,23,36]
                for column,mm in zip(t.columns,colwidth):column.width=Mm(mm)
                for row,values in zip(t.rows,v):
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for cell,item,mm in zip(row.cells,values,colwidth):
                        cell.width=Mm(mm);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);runs(p,item)
                        for r in p.runs:r.font.size=Pt(8.5 if width>7 else 9)
                        pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
                        for side in ['top','bottom','left','right']:
                            x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
                        pr.append(b)
                for c in t.rows[0].cells:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF1F2');c._tc.get_or_add_tcPr().append(sh)
                    for r in c.paragraphs[0].runs:r.bold=True
                repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
                d.add_paragraph().paragraph_format.space_after=Pt(1)
        path=a.out/f'manuscript_{mode}.docx';d.save(path);receipt[mode+'_sha256']=sha(path)
    (a.out/'assembly.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
if __name__=='__main__':main()
