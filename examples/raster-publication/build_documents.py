"""Typeset authored DEMO Markdown; do not generate or infer research claims."""
from pathlib import Path
import argparse,re,difflib,json,hashlib,copy
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Mm
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def blocks(path):
 lines=Path(path).read_text(encoding='utf-8').splitlines();out=[];i=0
 while i<len(lines):
  s=lines[i].strip()
  if not s:i+=1;continue
  if s.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in cells):rows.append(cells)
    i+=1
   out.append(('table',rows));continue
  if s.startswith('#'):
   n=len(s)-len(s.lstrip('#'));out.append(('title' if n==1 else 'h'+str(n),s[n:].strip()));i+=1;continue
  ls=[s.lstrip('> ').strip()];i+=1
  while i<len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(('#','|')):
   ls.append(lines[i].strip().lstrip('> '));i+=1
  out.append(('p',' '.join(ls)))
 return out

def clean_markup(t):return t.replace('**','').strip('*')
def add_runs(p,text,before=None,yellow=False):
 text=clean_markup(text)
 if not yellow or text==before:p.add_run(text);return
 old=before or '';new_tokens=re.findall(r'\s+|\S+',text);old_tokens=re.findall(r'\s+|\S+',old)
 for tag,a,b,c,d in difflib.SequenceMatcher(None,old_tokens,new_tokens,autojunk=False).get_opcodes():
  if c==d:continue
  r=p.add_run(''.join(new_tokens[c:d]))
  if tag!='equal':r.font.highlight_color=WD_COLOR_INDEX.YELLOW

def base_doc():
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=Inches(.75);s.top_margin=s.bottom_margin=Inches(.7)
 for name,size in [('Normal',11),('Title',18),('Heading 1',12),('Heading 2',11)]:
  st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
  for border in st._element.xpath('.//w:pBdr'):border.getparent().remove(border)
 for name in ['Heading 1','Heading 2']:d.styles[name].paragraph_format.keep_with_next=True
 hp=s.header.paragraphs[0];hp.add_run('ORIGINAL DEMO  |  SYNTHETIC VALUES  |  NOT EXPERIMENTAL RESULTS').font.size=Pt(8)
 return d

def write_doc(source,baseline,out,yellow=False,fig=None):
 items=blocks(source);old=[clean_markup(v) for k,v in blocks(baseline) if k!='table'];d=base_doc();mapping=[]
 for idx,(kind,value) in enumerate(items):
  if kind=='table':
   cols=len(value[0]);t=d.add_table(rows=1,cols=cols);t.autofit=False
   weights=[2.2,0.85]+[1]*(cols-2);widths=[Mm(177.8*w/sum(weights)) for w in weights]
   for j,w in enumerate(widths):t.columns[j].width=w
   for rownum,row in enumerate(value):
    cells=t.rows[0].cells if rownum==0 else t.add_row().cells
    assert len(row)==cols
    for j,val in enumerate(row):
     cells[j].width=widths[j];p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.05
     r=p.add_run(clean_markup(val));r.font.size=Pt(9);r.bold=rownum==0
     tcPr=cells[j]._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
     for side in ['top','left','bottom','right']:
      e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
     tcPr.append(b)
     if rownum==0:
      sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF0');tcPr.append(sh)
    trPr=t.rows[rownum]._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trPr.append(nosplit)
    if rownum==0:trPr.append(OxmlElement('w:tblHeader'))
   continue
  style='Title' if kind=='title' else 'Heading 1' if kind=='h2' else 'Heading 2' if kind.startswith('h') else 'Normal'
  p=d.add_paragraph(style=style);txt=clean_markup(value)
  candidate=max(old,key=lambda x:difflib.SequenceMatcher(None,x,txt).ratio(),default='')
  ratio=difflib.SequenceMatcher(None,candidate,txt).ratio();before=candidate if ratio>.4 else ''
  add_runs(p,txt,before,yellow)
  if before!=txt:mapping.append({'block':idx,'after':txt,'comparison_text':before,'comparison_scope':'first complete generation, not original research draft','marking':'new/replaced words yellow' if yellow else 'clean'})
 d.save(out);return {'source':str(source),'baseline':str(baseline),'output':str(out),'source_sha256':hashlib.sha256(Path(source).read_bytes()).hexdigest(),'output_sha256':hashlib.sha256(Path(out).read_bytes()).hexdigest(),'blocks':len(items),'tables':sum(k=='table' for k,v in items),'changes':mapping,'native_track_changes':False}

def figure_doc(image,caption,out,title='Original geometry demonstration'):
 d=base_doc();d.add_paragraph(title,'Title');d.add_picture(str(image),width=Mm(160));p=d.add_paragraph(Path(caption).read_text(encoding='utf-8'));p.paragraph_format.space_before=Pt(7)
 for r in p.runs:r.font.size=Pt(10)
 d.save(out)

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--source',required=True,type=Path);p.add_argument('--baseline',required=True,type=Path);p.add_argument('--out',required=True,type=Path);a=p.parse_args();a.out.mkdir(parents=True,exist_ok=False)
 results=[write_doc(a.source,a.baseline,a.out/'clean.docx'),write_doc(a.source,a.baseline,a.out/'review.docx',True)]
 (a.out/'build.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8');print('authored documents written')
